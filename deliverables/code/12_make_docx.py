# -*- coding: utf-8 -*-
"""12_make_docx.py —— A10：论文.md → 论文.docx（python-docx；pandoc 不可用）。
支持：标题层级/段落/有序无序列表/**粗体**渲染/Markdown 表格/图1~图8 按节插入。
产出：paper/论文.docx, output/logs/12_make_docx.log
"""
import sys, io, re
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ---------- LaTeX → Unicode 兜底转写层（主转写已在 14_md_unicode.py 完成于 md 源） ----------
TEX_MAP = [
    (r'\le', '≤'), (r'\ge', '≥'), (r'\leq', '≤'), (r'\geq', '≥'),
    (r'\in', '∈'), (r'\times', '×'), (r'\propto', '∝'), (r'\approx', '≈'),
    (r'\hat p', 'p̂'), (r'\hat u', 'û'), (r'\tilde q', 'q̃'),
    (r'\frac1n', '(1/n)'), (r'\frac{1}{n}', '(1/n)'),
    (r'\tau_1', 'τ₁'), (r'\tau_2', 'τ₂'), (r'\tau', 'τ'),
    (r'\alpha', 'α'), (r'\beta', 'β'), (r'\gamma', 'γ'), (r'\pi', 'π'),
    (r'\sum', 'Σ'), (r'\log', 'log'), (r'\cdot', '·'),
    (r'\{', '{'), (r'\}', '}'), (r'\_', '_'),
    (r'\,', ' '), (r'\;', ' '), (r'\qquad', '　'), (r'\quad', '　'),
    (r'\text{', ''), (r'\mathrm{', ''), (r'\begin{cases}', ''), (r'\end{cases}', ''),
    ('_1', '₁'), ('_2', '₂'), ('_i', 'ᵢ'), ('_k', 'ₖ'),
]

def texify(text):
    """兜底：把残余 LaTeX 模式转写为 Unicode 文本。"""
    if '$' not in text and '\\' not in text:
        return text
    text = text.replace('$$', '')
    text = re.sub(r'\$\$?([^$]*)\$\$?', r'\1', text)  # 行内 $...$ 剥壳
    for old, new in TEX_MAP:
        text = text.replace(old, new)
    text = text.replace('$', '')
    return text

def assert_no_latex(doc):
    """转换层单元自检：docx 全文不得含裸 $、反斜杠命令、begin{。"""
    bad = []
    for p in doc.paragraphs:
        t = p.text
        if '$' in t or re.search(r'\\[a-zA-Z]+', t) or 'begin{' in t:
            bad.append(t[:80])
    for tb in doc.tables:
        for row in tb.rows:
            for c in row.cells:
                t = c.text
                if '$' in t or re.search(r'\\[a-zA-Z]+', t) or 'begin{' in t:
                    bad.append('[TBL]' + t[:80])
    return bad

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
WS = r"C:\Users\21732\Desktop\2025b论文\agent_workspace_B1"
MD = rf"{WS}\paper\论文.md"
OUT = rf"{WS}\paper\论文.docx"
FIG = rf"{WS}\output\figures"
LOG = rf"{WS}\output\logs\12_make_docx.log"
lines = []
def log(msg=""):
    print(msg); lines.append(str(msg))

# 图插入映射：节标题关键词 → 图文件列表
FIG_MAP = {
    '4.4 附件1 标注结果与分析': ['fig1_标注规则二维散点.png', 'fig2_三类差额密度对比.png'],
    '4.5 敏感性分析': ['fig8_规则敏感性热图.png'],
    '5.3 结果分析': ['fig3_赔付金额分布对比.png', 'fig4_回归误差分析.png', 'fig7_特征消融.png'],
    '6.2': ['fig6_不均衡方案对比.png'],
    '6.3': ['fig5_两路线混淆矩阵.png'],
}

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_ea(run, ea='宋体', ascii_f='Times New Roman'):
    run.font.name = ascii_f
    run._element.rPr.rFonts.set(qn('w:eastAsia'), ea)

def add_para(text, size=10.5, bold=False, align=None, ea='宋体', space_after=6, indent=False):
    text = texify(text)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align:
        p.alignment = align
    # **bold** 分段渲染
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part)
            r.bold = bold
        r.font.size = Pt(size)
        set_ea(r, ea)
    return p

def add_heading_md(text, level):
    text = texify(text)
    sizes = {1: 16, 2: 14, 3: 12, 4: 11}
    p = doc.add_heading('', level=min(level, 4))
    r = p.add_run(text.replace('**', ''))
    r.font.size = Pt(sizes.get(min(level, 4), 11))
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.bold = True
    set_ea(r, '黑体')
    return p

def add_table_md(rows):
    rows = [texify(r) for r in rows]
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{2,}:?', c or '---') for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j in range(ncol):
            txt = row[j] if j < len(row) else ''
            cell = t.cell(i, j)
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(txt.replace('**', ''))
            r.font.size = Pt(9)
            r.bold = (i == 0)
            set_ea(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_fig(fname, caption, width=14.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(rf"{FIG}\{fname}", width=Cm(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.font.size = Pt(9); r.bold = True
    set_ea(r)

log("=== 12_make_docx start ===")
src = open(MD, encoding='utf-8').read()
lines_md = src.split('\n')

i = 0
n_tables = n_figs = 0
while i < len(lines_md):
    ln = lines_md[i]
    stripped = ln.strip()
    if not stripped:
        i += 1; continue
    if re.fullmatch(r'-{3,}|_{3,}|\*{3,}', stripped):  # Markdown 水平线：不写入 docx
        i += 1; continue
    # 表格块
    if stripped.startswith('|') and i + 1 < len(lines_md) and lines_md[i+1].strip().startswith('|'):
        block = []
        while i < len(lines_md) and lines_md[i].strip().startswith('|'):
            block.append(lines_md[i]); i += 1
        add_table_md(block); n_tables += 1
        continue
    # 标题
    m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
    if m:
        level, text = len(m.group(1)), m.group(2)
        if level == 1 and text.startswith('物流理赔'):
            add_para(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ea='黑体', space_after=12)
        else:
            add_heading_md(text, level)
        # 图插入
        for key, figs in FIG_MAP.items():
            if text.replace('：', ':').startswith(key) or key in text:
                for f in figs:
                    cap = f.split('_')[0].replace('fig', '图') + '  ' + f.split('_', 1)[1].replace('.png', '')
                    add_fig(f, cap)
                    n_figs += 1
        i += 1; continue
    # 列表
    m = re.match(r'^[-*]\s+(.*)$', stripped)
    if m:
        add_para('· ' + m.group(1), size=10.5, space_after=3)
        i += 1; continue
    m = re.match(r'^(\d+)[.．]\s+(.*)$', stripped)
    if m and len(m.group(1)) <= 2 and i > 20:  # 有序列表（避免把年份行误判）
        add_para(f"{m.group(1)}. {m.group(2)}", size=10.5, space_after=3)
        i += 1; continue
    # 普通段落
    add_para(stripped, indent=True)
    i += 1

# 转换层单元自检
bad = assert_no_latex(doc)
assert not bad, f"LaTeX 残留 {len(bad)} 处: {bad[:5]}"
log("转换层自检：全文无裸 $ / 无反斜杠命令 / 无 begin{} —— PASS")

doc.save(OUT)
log(f"[saved] {OUT}")
log(f"段落数={len(doc.paragraphs)} 表格数={n_tables} 插图数={n_figs}")
log("=== 12_make_docx done ===")
with open(LOG, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("\n[saved]", LOG)
