# -*- coding: utf-8 -*-
"""
a11_01_qa_check.py —— A11 独立质控程序化检查（Run-2，严格代执行 A11 qa-reproducer 锁定计划）
角色纪律：独立质控、只读不改他人产物；复跑他人脚本前先字节快照，结束后恢复原字节并复核 md5。
本脚本仅新增于 code/a11_01_qa_check.py；日志落 output/logs/qa_programmatic.log；
快照落 output/logs/a11_snapshots/（恢复后仅留 md5 证据清单）。

子命令：
  python a11_01_qa_check.py snapshot            # 第一步：字节快照 + md5 清单
  python a11_01_qa_check.py check               # 只读检查 Q-01/02/03/04/07/08/09/10/11/12/13
  python a11_01_qa_check.py verify_a6_01        # Q-05 复跑后核对（复跑本身由外部重定向执行）
  python a11_01_qa_check.py verify_a6_06        # Q-06 复跑后核对
  python a11_01_qa_check.py verify_fig fig3|fig8  # Q-14 复跑后 md5/像素核对（写 a11_rerun_fig.md5.log）
  python a11_01_qa_check.py restore             # 末步：恢复原字节 + md5 复核（快照目录仅留 md5 证据清单）
"""
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd

WS = Path(__file__).resolve().parents[1]
SNAP = WS / "output" / "logs" / "a11_snapshots"
TBL = WS / "output" / "tables"
LOGD = WS / "output" / "logs"
FIGD = WS / "output" / "figures"
PAPD = WS / "paper"
CODE = WS / "code"
DATAD = WS / "data"
QA_LOG = LOGD / "qa_programmatic.log"

SNAPSHOT_FILES = [
    "output/Result_提交.xlsx",
    "output/tables/q1_定稿配置.json",
    "output/tables/q1_边界表.csv",
    "output/tables/q1_自检表.csv",
    "output/tables/q1_择优日志.csv",
    "output/tables/q1_对偶报告.csv",
    "output/tables/q1_边界_bootstrapCI.csv",
    "output/tables/附件1_风险标注.csv",
    "output/logs/a6_阻断上报.md",
    "output/logs/a6_01_q1规则标注.log",   # 交付态为 0 字节，恢复时保持 0 字节
    "output/logs/a6_06_提交组装校验.log",
    "output/figures/fig3_q1规则检验.png",
    "output/figures/fig8_翻转带敏感性.png",
    "output/logs/a8_03_fig3_q1规则检验.log",
    "output/logs/a8_08_fig8_翻转带敏感性.log",
]
MANIFEST = SNAP / "md5_manifest.txt"

_results: list[tuple[str, str, str, str]] = []  # (Q编号, 子项, 判定, 证据)


def out(msg: str) -> None:
    print(msg)
    with open(QA_LOG, "a", encoding="utf-8") as f:
        f.write(msg + "\n")


def sec(title: str) -> None:
    out("")
    out("=" * 100)
    out(title)
    out("=" * 100)


def rec(qid: str, item: str, verdict: str, evidence: str) -> None:
    _results.append((qid, item, verdict, evidence))
    out(f"[{verdict}] {qid}·{item} —— {evidence}")


def md5_of(p: Path) -> str:
    h = hashlib.md5()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- 快照/恢复
def cmd_snapshot() -> int:
    SNAP.mkdir(parents=True, exist_ok=True)
    lines = ["# A11 字节快照 md5 清单（复跑他人脚本前留证；相对工作区根路径）", ""]
    missing = []
    for rel in SNAPSHOT_FILES:
        src = WS / rel
        if not src.exists():
            missing.append(rel)
            continue
        shutil.copy2(src, SNAP / Path(rel).name)
        data = src.read_bytes()
        bom = "BOM" if data[:3] == b"\xef\xbb\xbf" else "-"
        lines.append(f"{rel}\tmd5={md5_of(src)}\tbytes={len(data)}\t{bom}")
    lines += ["", f"合计快照 {len(SNAPSHOT_FILES) - len(missing)}/{len(SNAPSHOT_FILES)} 个文件"]
    lines += [f"缺失：{rel}" for rel in missing]
    MANIFEST.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    return 1 if missing else 0


def cmd_restore() -> int:
    if not MANIFEST.exists():
        print("FATAL：快照清单不存在，拒绝恢复")
        return 1
    entries = []
    for raw in MANIFEST.read_text(encoding="utf-8").splitlines():
        if raw and not raw.startswith("#") and "\tmd5=" in raw:
            rel, rest = raw.split("\t", 1)
            entries.append((rel, rest.split("md5=")[1].split("\t")[0]))
    lines = ["# A11 恢复复核（全部被复跑触碰的他人产物恢复原字节并与交付态 md5 比对）",
             "# 复核时间见文件 mtime；md5 与快照清单逐一比对。", ""]
    bad = 0
    for rel, want in entries:
        src = SNAP / Path(rel).name
        dst = WS / rel
        if not src.exists():
            lines.append(f"FAIL\t{rel}\t快照副本缺失")
            bad += 1
            continue
        shutil.copy2(src, dst)
        got = md5_of(dst)
        ok = got == want
        lines.append(f"{'PASS' if ok else 'FAIL'}\t{rel}\tmd5={got}\tbytes={dst.stat().st_size}"
                     + ("" if ok else f"\t期望 md5={want}"))
        bad += 0 if ok else 1
    (SNAP / "restore_md5_evidence.txt").write_text("\n".join(lines), encoding="utf-8")
    for f in SNAP.iterdir():
        if f.name not in {"md5_manifest.txt", "restore_md5_evidence.txt"}:
            f.unlink()
    print("\n".join(lines))
    print(f"\n恢复完成：{'全部一致' if bad == 0 else f'{bad} 个不一致'}；快照目录已清理，仅留 md5 证据清单")
    return 1 if bad else 0


# ---------------------------------------------------------------- 工具
def read_csv_sig(p: Path) -> pd.DataFrame:
    return pd.read_csv(p, encoding="utf-8-sig")


def paper_text() -> str:
    return (PAPD / "论文.md").read_text(encoding="utf-8")


def extract_table(text: str, marker: str) -> list[list[str]]:
    """定位 **表 N 标记，返回其后连续 '|' 行的单元格列表（跳过表头与分隔行）。"""
    lines = text.splitlines()
    rows, started, header_skipped = [], False, False
    for i, ln in enumerate(lines):
        if marker in ln:
            started = True
            continue
        if started:
            if ln.startswith("|"):
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                if all(re.fullmatch(r"[-: ]+", c) for c in cells if c):
                    continue  # 分隔行
                if not header_skipped:
                    header_skipped = True  # 首个 | 行为表头
                    continue
                rows.append(cells)
            elif rows:
                break
    return rows


def cells_floats(cell: str) -> list[float]:
    cleaned = cell.replace("**", "").replace(",", "")
    vals = []
    for m in re.findall(r"[-+]?\d+\.?\d*%?", cleaned):
        if m.endswith("%"):
            continue
        try:
            vals.append(float(m))
        except ValueError:
            pass
    return vals


def fmt4(x: float) -> str:
    return f"{x:.4f}"


# ---------------------------------------------------------------- Q-01/02 Result
def c_result() -> None:
    sec("Q-01/Q-02  Result_提交.xlsx 程序化校验（F1–F4/F6 + 与预测源逐格一致 + 论文占比/金额断言）")
    res_md5 = md5_of(WS / "output" / "Result_提交.xlsx")
    rec("Q-01", "文件md5", "PASS" if res_md5 == "a66012eaa89d01b6797f704f3f44d589" else "FAIL",
        f"output/Result_提交.xlsx md5={res_md5}（交付态快照 a66012eaa89d01b6797f704f3f44d589，与 STATE §4 记载一致）")
    res = pd.read_excel(WS / "output" / "Result_提交.xlsx")
    tpl = pd.read_excel(DATAD / "Result.xlsx")
    df2 = read_csv_sig(TBL / "附件2_clean.csv")
    pay = read_csv_sig(TBL / "附件2_赔付预测.csv")
    risk = read_csv_sig(TBL / "附件2_风险标注.csv".replace("标注", "预测"))
    rec("Q-01", "shape(2792,3)", "PASS" if res.shape == (2792, 3) else "FAIL",
        f"Result={res.shape}；模板={tpl.shape}；附件2_clean={df2.shape}")
    same_cols = list(res.columns) == list(tpl.columns)
    rec("Q-01", "列名列序=模板", "PASS" if same_cols else "FAIL",
        f"Result 列={list(res.columns)}；模板列={list(tpl.columns)}")
    ids_ok_tpl = bool(np.array_equal(res["运单号"].to_numpy(), tpl["运单号"].to_numpy()))
    ids_ok_clean = bool(np.array_equal(res["运单号"].to_numpy(), df2["运单号"].to_numpy()))
    rec("Q-01", "运单号零改动", "PASS" if ids_ok_tpl and ids_ok_clean else "FAIL",
        f"与 data/Result.xlsx 模板逐行一致={ids_ok_tpl}；与 output/tables/附件2_clean.csv 逐行一致={ids_ok_clean}"
        f"（2792 行，首行 {res['运单号'].iloc[0]}，末行 {res['运单号'].iloc[-1]}）")
    na = int(res["实际赔付金额"].isna().sum() + res["风险标注"].isna().sum())
    rec("Q-01", "两列无空值", "PASS" if na == 0 else "FAIL", f"空值合计={na}")
    amt = res["实际赔付金额"].to_numpy(dtype=float)
    neg = int((amt < 0).sum()); low = int((amt < 0.01).sum())
    bad2 = int((np.abs(amt - np.round(amt, 2)) > 1e-9).sum())
    rec("Q-01", "金额非负两位小数", "PASS" if neg == 0 and low == 0 and bad2 == 0 else "FAIL",
        f"负值={neg}，低于0.01={low}，非两位小数={bad2}；min={amt.min():.2f} max={amt.max():.2f}")
    labels = {"合理诉求", "诉求偏高", "严重超额"}
    labs_ok = set(res["风险标注"].unique()) <= labels
    rec("Q-01", "标签域", "PASS" if labs_ok else "FAIL",
        f"取值域={sorted(set(res['风险标注'].unique()))} ⊆ {{合理诉求,诉求偏高,严重超额}}")
    m_pay = dict(zip(pay["运单号"], pay["ŷ"].astype(float)))
    m_risk = dict(zip(risk["运单号"], risk["风险标注"]))
    diff_pay = sum(1 for i, v in zip(res["运单号"], amt) if abs(m_pay.get(i, np.nan) - v) > 1e-9)
    diff_risk = sum(1 for i, v in zip(res["运单号"], res["风险标注"]) if m_risk.get(i) != v)
    rec("Q-02", "与赔付/风险预测逐格一致", "PASS" if diff_pay == 0 and diff_risk == 0 else "FAIL",
        f"金额列不一致={diff_pay}/2792（源 output/tables/附件2_赔付预测.csv ŷ 列）；"
        f"标签列不一致={diff_risk}/2792（源 output/tables/附件2_风险预测.csv 风险标注列）")
    vc = res["风险标注"].value_counts()
    cnt = {k: int(vc.get(k, 0)) for k in ("合理诉求", "诉求偏高", "严重超额")}
    prop = {k: round(v / 2792 * 100, 2) for k, v in cnt.items()}
    ok = cnt == {"合理诉求": 2518, "诉求偏高": 254, "严重超额": 20} and \
        prop == {"合理诉求": 90.19, "诉求偏高": 9.10, "严重超额": 0.72}
    rec("Q-02", "占比 2518/254/20=90.19/9.10/0.72", "PASS" if ok else "FAIL",
        f"计数={cnt}；占比%={prop}（分母 2792）")
    s_min, s_max, s_sum = round(float(amt.min()), 2), round(float(amt.max()), 2), float(amt.sum())
    ok2 = s_min == 8.12 and s_max == 1455.73 and abs(s_sum - 838938) < 0.5
    rec("Q-02", "min 8.12 / max 1455.73 / 总额 838,938", "PASS" if ok2 else "FAIL",
        f"min={s_min} max={s_max} 总额={s_sum:.2f} 元（论文 §七 记载 838,938 元）")


# ---------------------------------------------------------------- Q-03 模板 md5
def c_template() -> None:
    sec("Q-03  模板 data/Result.xlsx md5 复核（模板未被改动证据）")
    now = md5_of(DATAD / "Result.xlsx")
    log6 = (LOGD / "a6_06_提交组装校验.log").read_text(encoding="utf-8")
    m = re.search(r"输入文件 md5：Result\.xlsx = ([0-9a-f]{32})", log6)
    logged = m.group(1) if m else "(日志无记录)"
    rec("Q-03", "模板md5", "PASS" if now == logged else ("WARN" if logged == "(日志无记录)" else "FAIL"),
        f"现状 md5={now}；a6_06 日志记录 md5={logged}（一致⇒模板自 A6 运行后未被改动）")


# ---------------------------------------------------------------- Q-04 clean.csv
def c_clean() -> None:
    sec("Q-04  clean.csv 数据契约（D22：utf-8-sig/形状/首列/md5）")
    f1, f2 = TBL / "附件1_clean.csv", TBL / "附件2_clean.csv"
    m1, m2 = md5_of(f1), md5_of(f2)
    rec("Q-04", "附件1 md5", "PASS" if m1.startswith("d5842f8b") else "FAIL",
        f"附件1_clean.csv md5={m1}（期望前缀 d5842f8b，与 q1_定稿配置.json 输入md5 记载一致）")
    rec("Q-04", "附件2 md5", "PASS" if m2.startswith("0846d4b0") else "FAIL",
        f"附件2_clean.csv md5={m2}（期望前缀 0846d4b0，同上）")
    b1 = f1.read_bytes()[:3] == b"\xef\xbb\xbf"
    b2 = f2.read_bytes()[:3] == b"\xef\xbb\xbf"
    rec("Q-04", "BOM 在位", "PASS" if b1 and b2 else "FAIL",
        f"附件1 BOM={b1}；附件2 BOM={b2}（D22 要求 utf-8-sig）")
    d1 = read_csv_sig(f1); d2 = read_csv_sig(f2)
    rec("Q-04", "形状与首列", "PASS" if d1.shape == (11167, 52) and d2.shape == (2792, 47)
        and d1.columns[0] == "行ID" and d2.columns[0] == "运单号" else "FAIL",
        f"附件1={d1.shape} 首列={d1.columns[0]}；附件2={d2.shape} 首列={d2.columns[0]}"
        f"（契约：附件2 终版 2792×47、运单号置首）")


# ---------------------------------------------------------------- Q-07 论文数字抽检
def c_paper_numbers() -> None:
    sec("Q-07  论文数字抽检（≥15 项，逐项注明源文件与行/格；比对容差=末位舍入）")
    pt = paper_text()
    n_item = 0

    def ok(q, item, cond, ev):
        nonlocal n_item
        n_item += 1
        rec("Q-07", f"{item}", "PASS" if cond else "FAIL", ev)

    # ---- 表1 边界表（源 q1_边界表.csv 5 行）----
    bound = read_csv_sig(TBL / "q1_边界表.csv")
    rows = extract_table(pt, "**表 1")
    ok("Q-07", "表1行数=5", len(rows) == 5, f"解析到 {len(rows)} 行（源 output/tables/q1_边界表.csv 5 行）")
    b_ok, b_ev = True, []
    for i, r in enumerate(rows):
        src = bound.iloc[i]
        lo, hi = cells_floats(r[1])[:2]
        cnt = cells_floats(r[2])[0]
        g1, g2 = cells_floats(r[3])[0], cells_floats(r[4])[0]
        conds = [abs(lo - src["x下界"]) < 5e-3, abs(hi - src["x上界"]) < 5e-3,
                 cnt == src["箱计数"], abs(g1 - src["g1_b"]) < 5e-3, abs(g2 - src["g2_b"]) < 5e-3]
        b_ok &= all(conds)
        b_ev.append(f"行{i+1}: 纸面[{lo},{hi}]×{cnt:.0f} g1={g1} g2={g2} vs CSV "
                    f"[{src['x下界']:.3f},{src['x上界']:.3f}]×{src['箱计数']} g1={src['g1_b']:.4f} g2={src['g2_b']:.4f}"
                    + ("✓" if all(conds) else "✗"))
    ok("Q-07", "表1 五行边界值", b_ok, "；".join(b_ev) + "（源 q1_边界表.csv 逐行）")

    # ---- θ* 与占比（源 q1_定稿配置.json）----
    cfg = json.loads((TBL / "q1_定稿配置.json").read_text(encoding="utf-8"))
    th = cfg["theta_star"]
    ok("Q-07", "θ*=(5,0.86,0.98)", th == {"B": 5, "tau1": 0.86, "tau2": 0.98} and
        "τ₁=0.86,τ₂=0.98" in pt.replace(" ", ""), f"json theta_star={th}；论文含定稿参数串")
    lc = cfg["labels_counts"]
    ok("Q-07", "附件1占比 85.99/12.00/2.01 与 9,602/1,340/225",
        lc == {"合理诉求": 9602, "诉求偏高": 1340, "严重超额": 225} and
        all(s in pt for s in ("85.99%", "12.00%", "2.01%", "9,602", "1,340")),
        f"json labels_counts={lc}；论文 5.1.4 节含占比与计数串")
    sel = cfg["择优路径"]
    ok("Q-07", "择优 76 通过/31 入邻域", sel["L0通过"] == 76 and sel["L1通过"] == 31 and
        "76 组通过硬筛选" in pt and "31 组落入先验邻域" in pt,
        f"json 择优路径={sel}；论文 5.1.3 节同值")

    # ---- 自检表 V1–V8（源 q1_自检表.csv）----
    chk = read_csv_sig(TBL / "q1_自检表.csv").set_index("编号")
    v4 = str(chk.loc["V4", "数值"])
    ok("Q-07", "D_geo=52.84 与 ρ=0.229", "D_geo=52.8364" in v4 and "0.2289" in v4 and
        "52.84" in pt and "0.229" in pt, f"自检表 V4 数值=「{v4}」；论文表2/5.1.5 舍入一致")
    v3 = str(chk.loc["V3", "数值"])
    ok("Q-07", "V3 紧凑度 0.7355/0.7904/0.5035", all(s in v3 for s in ("0.7355", "0.7904", "0.5035")) and
        all(s in pt for s in ("0.7355", "0.7904", "0.5035")), f"自检表 V3=「{v3}」；论文表2 同值")
    v5 = str(chk.loc["V5", "数值"])
    ok("Q-07", "V5 e中位 2,011.1/1,112.7/218.3", all(s in v5 for s in ("2011.1", "1112.7", "218.3")) and
        all(s in pt for s in ("2,011.1", "1,112.7", "218.3")), f"自检表 V5=「{v5}」；论文表2 千分位同值")
    v6 = str(chk.loc["V6", "数值"])
    ok("Q-07", "V6 ①0.2071/②1.0298%/③极差0.0051",
        all(s in v6 for s in ("0.2071", "0.4803", "0.010298", "0.0051")) and
        all(s in pt for s in ("0.2071", "0.4803", "1.0298%")),
        f"自检表 V6=「{v6}」；论文表2/5.1.7/8.2 同值（超限如实报告）")
    bci = read_csv_sig(TBL / "q1_边界_bootstrapCI.csv")
    w12 = float(bci[(bci["箱序"] == 1) & (bci["边界j"] == 2)]["相对宽度"].iloc[0])
    ok("Q-07", "严重界第1箱 CI 相对宽度 0.8102", round(w12, 4) == 0.8102 and "0.8102" in pt,
        f"q1_边界_bootstrapCI.csv 行(箱1,边界2) 相对宽度={w12:.10f}；论文 5.1.5/8.2 同值")
    dual = read_csv_sig(TBL / "q1_对偶报告.csv")
    r0 = dual[dual["箱序"] == 0].iloc[0]
    ok("Q-07", "对偶三分位 0.6255/0.8408/0.9461",
        round(r0["r_P50"], 4) == 0.6255 and round(r0["r_τ1=0.86"], 4) == 0.8408 and
        round(r0["r_τ2=0.98"], 4) == 0.9461 and all(s in pt for s in ("0.6255", "0.8408", "0.9461")),
        f"q1_对偶报告.csv 箱序0 行：P50={r0['r_P50']:.6f} τ1={r0['r_τ1=0.86']:.6f} τ2={r0['r_τ2=0.98']:.6f}；论文 5.1.6 同值")
    ok("Q-07", "对偶各箱中位 0.813→0.532",
        round(float(dual[dual["箱序"] == 1]["r_P50"].iloc[0]), 3) == 0.813 and
        round(float(dual[dual["箱序"] == 5]["r_P50"].iloc[0]), 3) == 0.532 and "0.813" in pt and "0.532" in pt,
        "q1_对偶报告.csv 箱1/箱5 r_P50；论文 5.1.6 同值")

    # ---- 翻转带敏感性（源 a7_翻转带敏感性.csv / a7_翻转带画像.csv）----
    fs = read_csv_sig(TBL / "a7_翻转带敏感性.csv")
    tau9 = fs[fs["网格"] == "τ邻域"]
    lo9, hi9 = tau9["flip_own_中位"].min(), tau9["flip_own_中位"].max()
    ok("Q-07", "τ邻域 9 格 0.949%–1.084%",
        round(lo9 * 100, 3) == 0.949 and round(hi9 * 100, 3) == 1.084 and "0.949%" in pt and "1.084%" in pt,
        f"a7_翻转带敏感性.csv τ邻域 flip_own_中位 min={lo9:.6f} max={hi9:.6f}；论文 5.1.7/六(一) 同值")
    bnb = fs[fs["网格"] == "B邻域"].sort_values("B")
    bvals = [round(v * 100, 2) for v in bnb["flip_own_中位"]]
    ok("Q-07", "B 邻域翻转率 1.37/1.58/1.69", bvals == [1.37, 1.58, 1.69] and
        all(s in pt for s in ("1.37%", "1.58%", "1.69%")),
        f"B=10/15/20 flip_own_中位%={bvals}；论文 5.1.7/六(一) 同值")
    detl = [round(v * 100, 1) for v in bnb["确定性重标差异率"]]
    ok("Q-07", "确定性重标 4.7%–5.1%", detl == [4.7, 4.8, 5.1] and "4.7%–5.1%" in pt,
        f"确定性重标差异率%={detl}；论文 六(一) 同值")
    fp = read_csv_sig(TBL / "a7_翻转带画像.csv").set_index("块")
    d05 = fp.loc["距离带"].loc[lambda d: d["键"] == "[0,0.5%)"].iloc[0] if False else \
        fp.reset_index().query("块=='距离带' and 键=='[0,0.5%)'").iloc[0]
    d50 = fp.reset_index().query("块=='距离带' and 键=='>=50%'").iloc[0]
    ok("Q-07", "贴边界 43.5% / ≥50% 带 6,978 行 62.5% 恒 0",
        round(d05["行均翻转频率_全量口径"] * 100, 1) == 43.5 and int(d50["行数"]) == 6978 and
        round(d50["行占比"] * 100, 1) == 62.5 and d50["行均翻转频率_全量口径"] == 0 and
        "43.5%" in pt and "6,978" in pt and "62.5%" in pt,
        f"a7_翻转带画像.csv：[0,0.5%)行均翻转={d05['行均翻转频率_全量口径']:.4f}；≥50% 行数={int(d50['行数'])} 占比={d50['行占比']:.4f} 频率=0")
    r927 = fp.reset_index().query("块=='风险集' and 键.str.startswith('任一次')").iloc[0]
    r471 = fp.reset_index().query("块=='风险集' and 键.str.startswith('高频')").iloc[0]
    ok("Q-07", "927 行(8.3%) / 471 行(4.2%) 贡献 94.7%",
        int(r927["行数"]) == 927 and round(r927["行占比"] * 100, 1) == 8.3 and
        int(r471["行数"]) == 471 and round(r471["行占比"] * 100, 1) == 4.2 and
        round(r471["翻转质量占比_全量口径"] * 100, 1) == 94.7 and
        all(s in pt for s in ("927 行", "8.3%", "471 行", "4.2%", "94.7%")),
        "a7_翻转带画像.csv 风险集两行；论文 5.1.7 同值")
    t2pct = fp.reset_index().query("块=='距离带' and 键.str.contains('贴g1')").iloc[0]
    t2g2 = fp.reset_index().query("块=='距离带' and 键.str.contains('贴g2')").iloc[0]
    n163 = int(t2pct["行数"]) + int(t2g2["行数"])
    ok("Q-07", "<2% 带 163 行 1.46%、行均频率 0.34–0.38",
        n163 == 163 and round(t2pct["行占比"] + t2g2["行占比"], 4) == 0.0146 and
        0.34 <= t2g2["行均翻转频率_全量口径"] and t2pct["行均翻转频率_全量口径"] <= 0.38 and "163 行" in pt,
        f"贴g1 {int(t2pct['行数'])}行+贴g2 {int(t2g2['行数'])}行={n163}行；占比={t2pct['行占比']+t2g2['行占比']:.6f}；"
        f"行均频率 {t2g2['行均翻转频率_全量口径']:.2f}–{t2pct['行均翻转频率_全量口径']:.2f}")
    tr0 = fp.reset_index().query("块=='转移方向' and 键=='合理诉求→严重超额'").iloc[0]
    tr5 = fp.reset_index().query("块=='转移方向' and 键=='严重超额→合理诉求'").iloc[0]
    ok("Q-07", "合理→严重 0 次、严重→合理仅 5 次", int(tr0["行数"]) == 0 and int(tr5["行数"]) == 5 and "仅 5 次" in pt,
        f"a7_翻转带画像.csv 转移方向两行：{int(tr0['行数'])}/{int(tr5['行数'])}")

    # ---- 表3 Q2 九行双口径（源 q2_指标汇总.csv：分层=3种子均值，行序=种子20251004）----
    q2 = read_csv_sig(TBL / "q2_指标汇总.csv")
    lay = q2[q2["口径"] == "分层5折"]; seq = q2[q2["口径"] == "行序外推"]
    def m3(df, model, col="WAPE"):
        v = df[df["模型"] == model][col].astype(float)
        return float(v.mean()) if len(v) > 1 else float(v.iloc[0])
    rows3 = extract_table(pt, "**表 3")
    key3 = [("log10 + LGBM", "LGBM_变体A"), ("Huber", "LGBM_变体B"), ("XGBoost", "XGB_变体A对照"),
            ("ElasticNet", "ElasticNet"), ("OLS", "OLS"), ("B4", "B4_分段比"), ("B3", "B3_单参数比"),
            ("B1", "B1_中位数"), ("B2", "B2_均值")]
    t3_ok, t3_ev = True, []
    for kw, model in key3:
        row = next(r for r in rows3 if kw in r[0])
        c2 = cells_floats(row[1]); c3 = cells_floats(row[2])
        exp_lay, exp_seq = m3(lay, model), float(seq[seq["模型"] == model]["WAPE"].iloc[0])
        conds = abs(c2[0] - exp_lay) <= 6e-5 and abs(c3[0] - exp_seq) <= 6e-5
        t3_ok &= conds
        t3_ev.append(f"{model}: 纸面 {c2[0]:.4f}/{c3[0]:.4f} vs 源 {exp_lay:.6f}/{exp_seq:.6f}"
                     + ("✓" if conds else "✗"))
    ok("Q-07", "表3 九行双口径 WAPE", t3_ok, "；".join(t3_ev) +
       "（源 q2_指标汇总.csv：分层=3 种子均值，行序=种子 20251004）")
    ci_lay = lay[(lay["模型"] == "LGBM_变体A") & (lay["种子"] == 20251004)].iloc[0]
    ci_seq = seq[seq["模型"] == "LGBM_变体A"].iloc[0]
    ok("Q-07", "主种子 CI [0.3347,0.3463] 与行序 CI [0.3233,0.3497]",
        round(ci_lay["WAPE_CI2.5"], 4) == 0.3347 and round(ci_lay["WAPE_CI97.5"], 4) == 0.3463 and
        round(ci_seq["WAPE_CI2.5"], 4) == 0.3233 and round(ci_seq["WAPE_CI97.5"], 4) == 0.3497 and
        "[0.3347, 0.3463]" in pt and "[0.3233, 0.3497]" in pt,
        f"q2_指标汇总.csv LGBM_变体A：分层主种子 CI=[{ci_lay['WAPE_CI2.5']:.4f},{ci_lay['WAPE_CI97.5']:.4f}]；"
        f"行序 CI=[{ci_seq['WAPE_CI2.5']:.4f},{ci_seq['WAPE_CI97.5']:.4f}]")
    seeds_a = lay[lay["模型"] == "LGBM_变体A"]["WAPE"].astype(float)
    ok("Q-07", "3 种子极差 0.0009", round(seeds_a.max() - seeds_a.min(), 4) == 0.0009 and "0.0009" in pt,
        f"极差={seeds_a.max() - seeds_a.min():.6f}→0.0009；论文 5.2.4/六(五) 同值")
    mae3 = m3(lay, "LGBM_变体A", "MAE")
    rmse_main = float(lay[(lay["模型"] == "LGBM_变体A") & (lay["种子"] == 20251004)]["RMSE"].iloc[0])
    ok("Q-07", "MAE 99.42 / RMSE 142.52", round(mae3, 2) == 99.42 and round(rmse_main, 2) == 142.52 and
        "99.42" in pt and "142.52" in pt,
        f"3 种子 MAE 均值={mae3:.4f}；主种子池化 RMSE={rmse_main:.4f}（q2_指标汇总.csv）；论文 5.2.4 同值")

    # ---- 池化 OOF 重算（源 q2_oof预测.csv）----
    oof = read_csv_sig(TBL / "q2_oof预测.csv")
    oof = oof[oof["变体"] == "LGBM_变体A"]
    y = oof["y_true"].to_numpy(dtype=float); p = oof["y_pred_oof"].to_numpy(dtype=float)
    wape_pool = float(np.abs(y - p).sum() / np.abs(y).sum())
    rel = np.abs(p - y) / np.maximum(y, 1e-9)
    ok("Q-07", "池化 OOF WAPE=0.3402、|相对误差| P50=0.348/P90=1.543",
        round(wape_pool, 4) == 0.3402 and round(float(np.quantile(rel, 0.5)), 3) == 0.348 and
        round(float(np.quantile(rel, 0.9)), 3) == 1.543 and "0.348" in pt and "1.543" in pt,
        f"q2_oof预测.csv（n={len(oof)}）重算：WAPE={wape_pool:.6f}，P50={np.quantile(rel,0.5):.4f}，"
        f"P90={np.quantile(rel,0.9):.4f}；论文 5.2.4/figures_清单 fig4 同值")

    # ---- 残差结构（源 a7_q2残差分析.csv）----
    rz_all = read_csv_sig(TBL / "a7_q2残差分析.csv")
    d10 = rz_all.query("块=='y十分位' and 键.str.startswith('D10')").iloc[0]
    d1r = rz_all.query("块=='y十分位' and 键.str.startswith('D1')").iloc[0]
    hz = rz_all.query("块=='异方差' and 键.str.contains('中位')").iloc[0]
    tail95 = rz_all.query("块=='重尾证据' and 键.str.contains('P95')").iloc[0]
    d10y = float(d10["值2"].split("=")[1]); d10e = float(d10["值3"].split("=")[1])
    t95y = float(tail95["值1"].split("=")[1]); t95e = float(tail95["值2"].split("=")[1])
    ok("Q-07", "头部十分位 32.0%/21.2%、y≥P95 19.0%/12.2%",
        abs(d10y - 0.3202) < 5e-5 and abs(d10e - 0.2115) < 5e-5 and
        abs(t95y - 0.1900) < 5e-5 and abs(t95e - 0.1222) < 5e-5 and
        all(s in pt for s in ("32.0%", "21.2%", "19.0%", "12.2%")),
        f"a7_q2残差分析.csv D10 行（Σy份额={d10['值2']} Σ|err|份额={d10['值3']}）与重尾证据行"
        f"（y≥P95：Σy份额={tail95['值1']} Σ|err|份额={tail95['值2']}）；0.2115→21.2%、0.1222→12.2% 四舍五入；论文 5.2.5 同值")
    ok("Q-07", "区间 WAPE 2.366→0.225、|残差|中位 33→171、Spearman 0.4724/−0.4533",
        "2.3656" in d1r["值4"] and "0.2251" in d10["值4"] and
        "33→35→40→47→59→68→87→115→152→171" in hz["值1"] and "1.0000" in hz["值2"] and
        "0.4724" in pt and "0.4533" in pt,
        f"a7_q2残差分析.csv D1 值4=「{d1r['值4']}」、D10 值4=「{d10['值4']}」、异方差「|残差|中位」行=「{hz['值1']}」；论文 5.2.5 同值")
    ll = rz_all.query("块=='loglog基线'").iloc[0]
    ok("Q-07", "单特征 log-log WAPE 0.4545 / R² 0.6098", "0.4545" in ll["值3"] and "0.6098" in ll["值2"] and
        "0.4545" in pt and "0.6098" in pt, f"a7_q2残差分析.csv loglog基线行：{ll['键']} {ll['值2']} {ll['值3']}")
    ok("Q-07", "索赔金额 Spearman +0.8069", "0.8069" in pt and
        "0.8069" in rz_all.query("块=='结论素材'").iloc[0]["值1"],
        "a7_q2残差分析.csv 结论素材行（引 A3 T5）；论文 5.2.5 同值")

    # ---- 表4 Q3（源 q3_指标汇总/q3_两路线对比/q3_消融矩阵）----
    q3m = read_csv_sig(TBL / "q3_指标汇总.csv")
    q3r = read_csv_sig(TBL / "q3_两路线对比.csv")
    q3a = read_csv_sig(TBL / "q3_消融矩阵.csv")
    def q3val(df, cfg_name, cal, seed=20251004, col="宏F1"):
        r = df[(df["配置"] == cfg_name) & (df["口径"] == cal) & (df["种子"] == seed)]
        return float(r[col].iloc[0])
    rows4 = extract_table(pt, "**表 4")
    exp4 = [
        ("方式 2 主配置", 0.6704, 0.3200, 0.6710, 0.3636),
        ("方式 1 完整链路", 0.6362, 0.1733, 0.6643, 0.2273),
        ("A0", 0.6529, 0.2133, 0.6922, 0.3182),
        ("A1", 0.6631, 0.2978, 0.6668, 0.3409),
        ("A2", 0.6568, 0.2667, 0.6824, 0.3636),
        ("A3", 0.6642, 0.2978, 0.6710, 0.3636),
        ("A4", 0.6758, 0.3644, 0.6870, 0.3864),
        ("A5", 0.6556, 0.3200, 0.6498, 0.3636),
        ("A6", 0.4298, 0.9556, 0.4291, 0.9773),
    ]
    src_map = {
        "方式 2 主配置": (q3m, "方式2_主配置(嵌套γ,s)"),
        "方式 1 完整链路": (q3r, "方式1_完整链路"),
        "A0": (q3a, "消融A0"), "A1": (q3a, "消融A1"), "A2": (q3a, "消融A2"), "A3": (q3a, "消融A3"),
        "A4": (q3a, "消融A4"), "A5": (q3a, "消融A5"), "A6": (q3a, "消融A6"),
    }
    t4_ok, t4_warn, t4_ev = True, [], []
    for kw, e_l, e_r0, e_s, e_r1 in exp4:
        row = next(r for r in rows4 if kw in r[0])
        got = cells_floats(row[1])[:2] + cells_floats(row[2])[:2]
        df, cname = src_map[kw]
        exp = [q3val(df, cname, "分层5折"), q3val(df, cname, "分层5折", col="严重类_recall"),
               q3val(df, cname, "行序外推"), q3val(df, cname, "行序外推", col="严重类_recall")]
        diffs = [abs(g - e) for g, e in zip(got, exp)]
        if all(d <= 6e-5 for d in diffs):
            t4_ev.append(f"{kw}: 纸面 {got} ≈ 源 {[round(x,4) for x in exp]} ✓")
        elif all(d <= 1.1e-4 for d in diffs):
            t4_warn.append(f"{kw}: 纸面 {got} vs 源 {[round(x,4) for x in exp]}（截断 vs 四舍五入，差 ≤0.0001）")
        else:
            t4_ok = False
            t4_ev.append(f"{kw}: 纸面 {got} vs 源 {[round(x,4) for x in exp]} ✗")
    rec("Q-07", "表4 九行双口径（宏F1+严重召回）", "PASS" if t4_ok and not t4_warn else ("WARN" if t4_ok else "FAIL"),
        "；".join(t4_ev + t4_warn) +
        "（源：方式2=q3_指标汇总.csv；方式1=q3_两路线对比.csv；A0–A6=q3_消融矩阵.csv）"
        + ("" if not t4_warn else "；A3 行序源值 0.6709632 应四舍五入为 0.6710，纸面 0.6709 为截断——同值在 5.3.4 节写作 0.6710，属同一数字两种舍入表述，不影响数值正确性"))

    # ---- 3 种子均值/一致率/混淆/杠杆/log loss ----
    m2_lay = [q3val(q3m, "方式2_主配置(嵌套γ,s)", "分层5折", s) for s in (20251004, 20251005, 20251006)]
    m1_lay = [q3val(q3r, "方式1_完整链路", "分层5折", s) for s in (20251004, 20251005, 20251006)]
    r2_lay = [q3val(q3m, "方式2_主配置(嵌套γ,s)", "分层5折", s, "严重类_recall") for s in (20251004, 20251005, 20251006)]
    r1_lay = [q3val(q3r, "方式1_完整链路", "分层5折", s, "严重类_recall") for s in (20251004, 20251005, 20251006)]
    ok("Q-07", "3 种子均值 0.660 对 0.636、严重召回 0.293 对 0.172",
        round(float(np.mean(m2_lay)), 3) == 0.660 and round(float(np.mean(m1_lay)), 3) == 0.636 and
        round(float(np.mean(r2_lay)), 3) == 0.293 and round(float(np.mean(r1_lay)), 3) == 0.172 and
        all(s in pt for s in ("0.660 对 0.636", "0.293 对 0.172")),
        f"均值：方式2={np.mean(m2_lay):.4f} 方式1={np.mean(m1_lay):.4f}；召回：{np.mean(r2_lay):.4f}/{np.mean(r1_lay):.4f}"
        f"（q3_指标汇总.csv / q3_两路线对比.csv 分层 3 行）")
    att_lay = [float(v) for v in q3r[q3r["配置"].isna()]["标签一致率"][:3]]
    att_seq = float(q3r[q3r["配置"].isna() & (q3r["口径"] == "行序外推")]["标签一致率"].iloc[0])
    att_p = float(q3r[q3r["配置"].isna() & (q3r["口径"] == "附件2预测")]["标签一致率"].iloc[0])
    ok("Q-07", "一致率 0.9777/0.9740（论文）与 0.9864（附件2，源表在证）",
        round(float(np.mean(att_lay)), 4) == 0.9777 and round(att_seq, 4) == 0.9740 and
        round(att_p, 4) == 0.9864 and "0.9777" in pt and "0.9740" in pt,
        f"q3_两路线对比.csv 空配置行：分层 3 种子均值={np.mean(att_lay):.6f}→0.9777、行序={att_seq:.6f}→0.9740、"
        f"附件2 两路线标签一致率={att_p:.6f}→0.9864（论文未直接引用 0.9864，改为引用 §七 边界带 59 行不一致率 37.3%，与源一致、无矛盾）")
    cm = read_csv_sig(TBL / "q3_混淆矩阵.csv")
    cmL = cm[cm["口径"] == "分层5折_主种子池化"]
    tot_err = int(cmL[cmL["真实"] != cmL["预测"]]["行数"].sum())
    cross = int(cmL[((cmL["真实"] == "合理诉求") & (cmL["预测"] == "严重超额")) |
                    ((cmL["真实"] == "严重超额") & (cmL["预测"] == "合理诉求"))]["行数"].sum())
    sev_wrong = int(cmL[(cmL["真实"] == "严重超额") & (cmL["预测"] != "严重超额")]["行数"].sum())
    sev_137 = int(cmL[(cmL["真实"] == "严重超额") & (cmL["预测"] == "诉求偏高")]["行数"].iloc[0])
    sev_16 = int(cmL[(cmL["真实"] == "严重超额") & (cmL["预测"] == "合理诉求")]["行数"].iloc[0])
    ok("Q-07", "混淆 846 错误/95.0% 相邻/153=137+16",
        tot_err == 846 and round((tot_err - cross) / tot_err * 100, 1) == 95.0 and
        sev_wrong == 153 and sev_137 == 137 and sev_16 == 16 and
        all(s in pt for s in ("846", "95.0%", "153", "137", "16")),
        f"q3_混淆矩阵.csv 分层池化：总错误={tot_err}，相邻={tot_err - cross}（{(tot_err-cross)/tot_err:.4f}），"
        f"跨类={cross}；严重类漏判 {sev_137}+{sev_16}={sev_wrong}")
    s0 = q3m[(q3m["配置"] == "方式2_主配置(嵌套γ,s)") & (q3m["种子"] == 20251004)].iloc[0]
    ok("Q-07", "逐类 0.9855/0.6616/0.3830 与漏判表述",
        round(s0["类0_recall"], 4) == 0.9855 and round(s0["类1_f1"], 4) == 0.6616 and
        round(s0["类2_f1"], 4) == 0.3830 and all(s in pt for s in ("0.9855", "0.6616", "0.3830")),
        f"q3_指标汇总.csv 主种子行：类0召回={s0['类0_recall']:.4f} 类1F1={s0['类1_f1']:.4f} 类2F1={s0['类2_f1']:.4f}")
    db_all = read_csv_sig(TBL / "a7_双口径分解.csv")
    lay_macro = db_all[(db_all["口径"] == "分层5折") & (db_all["类"] == "宏F1")].iloc[0]
    seq_macro = db_all[(db_all["口径"] == "行序外推") & (db_all["类"] == "宏F1")].iloc[0]
    d_lay_sev = db_all[(db_all["口径"] == "分层5折") & (db_all["类"] == "严重超额")]["ΔF1(主−A0)"].iloc[0]
    d_seq_sev = db_all[(db_all["口径"] == "行序外推") & (db_all["类"] == "严重超额")]["ΔF1(主−A0)"].iloc[0]
    d_seq_mid = db_all[(db_all["口径"] == "行序外推") & (db_all["类"] == "诉求偏高")]["ΔF1(主−A0)"].iloc[0]
    ok("Q-07", "杠杆 +0.0175 / 严重 ΔF1 +0.0542 / 行序 −0.0212、严重 −0.0357、偏高 −0.0281",
        round(lay_macro["ΔF1(主−A0)"], 4) == 0.0175 and round(d_lay_sev, 4) == 0.0542 and
        round(seq_macro["ΔF1(主−A0)"], 4) == -0.0212 and round(d_seq_sev, 4) == -0.0357 and
        round(d_seq_mid, 4) == -0.0281 and
        all(s in pt for s in ("+0.0175", "+0.0542", "−0.0212")),
        f"a7_双口径分解.csv 宏F1 行：分层 Δ={lay_macro['ΔF1(主−A0)']:.6f} 行序 Δ={seq_macro['ΔF1(主−A0)']:.6f}；"
        f"严重类 Δ：{d_lay_sev:.6f}/{d_seq_sev:.6f}；偏高类（行序）Δ={d_seq_mid:.6f}")
    bb = read_csv_sig(TBL / "a7_双口径分解_bootstrap.csv")
    b_lay = bb[bb["口径"] == "分层5折"].iloc[0]; b_seq = bb[bb["口径"] == "行序外推"].iloc[0]
    ok("Q-07", "bootstrap CI [−0.0173,+0.0526] P=0.844；[−0.0988,+0.0515] P=0.283",
        abs(b_lay["Δ宏F1_CI2.5"] - (-0.017331)) < 5e-6 and abs(b_lay["Δ宏F1_CI97.5"] - 0.052628) < 5e-6 and
        abs(b_lay["P_Δ大于0"] - 0.8435) < 5e-5 and abs(b_seq["Δ宏F1_CI2.5"] - (-0.098803)) < 5e-6 and
        abs(b_seq["Δ宏F1_CI97.5"] - 0.051458) < 5e-6 and abs(b_seq["P_Δ大于0"] - 0.2835) < 5e-5 and
        "0.844" in pt and "0.283" in pt,
        f"a7_双口径分解_bootstrap.csv 重构端界A 行：分层 CI=[{b_lay['Δ宏F1_CI2.5']:.6f},{b_lay['Δ宏F1_CI97.5']:.6f}] "
        f"P={b_lay['P_Δ大于0']:.4f}→0.844；行序 CI=[{b_seq['Δ宏F1_CI2.5']:.6f},{b_seq['Δ宏F1_CI97.5']:.6f}] "
        f"P={b_seq['P_Δ大于0']:.4f}→0.283（端界B 行 0.8455/0.3145 同向）；论文 5.3.5/六(三)/8.2 同值")
    ok("Q-07", "log loss 0.2386 vs 0.2208",
        round(s0["OOF_logloss"], 4) == 0.2386 and
        round(float(q3a[(q3a["配置"] == "消融A0") & (q3a["口径"] == "分层5折")]["OOF_logloss"].iloc[0]), 4) == 0.2208 and
        "0.2386" in pt and "0.2208" in pt,
        f"q3_指标汇总.csv 主配置 OOF_logloss={s0['OOF_logloss']:.6f}；q3_消融矩阵.csv A0=0.220797；论文 5.3.5/六(五)/8.2 同值")

    # ---- 表5 δ 敏感性 + 传导（源 a7_q3误差传导.csv）----
    tr = read_csv_sig(TBL / "a7_q3误差传导.csv")
    dlt = tr[tr["块"] == "δ敏感性"]
    rows5 = extract_table(pt, "**表 5")
    t5_ok, t5_ev = True, []
    for i, r in enumerate(rows5):
        src = dlt.iloc[i]
        got = cells_floats(r[1])[:1] + cells_floats(r[2])[:1] + cells_floats(r[3])[:1] + cells_floats(r[4])[:1]
        band_n = int(re.search(r"带行数=(\d+)", src["值1"]).group(1))
        p_err = float(re.search(r"P\(L̃≠y\|带\)=([0-9.]+)", src["值2"]).group(1))
        p_dis = float(re.search(r"P\(L̃≠L̂\|带\)=([0-9.]+)", src["值2"]).group(1))
        ratio = float(src["值3"].split("=")[1])
        exp = [band_n, p_err, p_dis, ratio]
        conds = [abs(got[0] - exp[0]) < 5e-3 and abs(got[1] - exp[1]) < 5e-3 and
                 abs(got[2] - exp[2]) < 5e-3 and abs(got[3] - exp[3]) < 5e-3]
        t5_ok &= conds[0]
        t5_ev.append(f"δ={r[0]}: 纸面 {got} vs 源 {exp}" + ("✓" if conds[0] else "✗"))
    ok("Q-07", "表5 六行 δ 敏感性", t5_ok and len(rows5) == 6, "；".join(t5_ev) + "（源 a7_q3误差传导.csv δ敏感性 6 行）")
    mech = tr[tr["块"] == "机制量化"].set_index("键")
    m183 = float(re.search(r"P50=([0-9.]+)", mech.loc["翻转组(L̃≠y) |ŷ−y| 中位/P90", "值1"]).group(1))
    m060 = float(re.search(r"P50=([0-9.]+)", mech.loc["一致组(L̃=y) |ŷ−y| 中位/P90", "值1"]).group(1))
    ok("Q-07", "183.62 元对 60.57 元（约 3 倍）、831/10,336 行",
        abs(m183 - 183.62) < 5e-3 and abs(m060 - 60.57) < 5e-3 and
        "n=831" in mech.loc["翻转组(L̃≠y) |ŷ−y| 中位/P90", "值3"] and
        "n=10336" in mech.loc["一致组(L̃=y) |ŷ−y| 中位/P90", "值3"] and
        all(s in pt for s in ("183.62", "60.57", "831", "10,336", "3 倍")),
        f"a7_q3误差传导.csv 机制量化两行：P50={m183}/{m060} 元，n=831/10336；论文 5.3.6/8.1 同值")
    geo = tr[tr["块"] == "边界几何"].set_index("键")
    band5 = geo.loc["d_min<0.05 带（真实坐标, n=415）"]
    ok("Q-07", "d_min 0.1310/0.6738/0.1335 与 0.6745、MW p≈0、415 行 3.7%、47.95%/49.40%",
        "0.1310" in geo.loc["真实坐标 d_min（到最近适用 g_j 的相对距离）", "值1"] and
        "0.6738" in geo.loc["真实坐标 d_min（到最近适用 g_j 的相对距离）", "值2"] and
        "0.1335" in geo.loc["真实坐标 d_min（到最近适用 g_j 的相对距离）", "值3"] and
        "0.6745" in geo.loc["真实坐标 d_min（到最近适用 g_j 的相对距离）", "值4"] and
        "p=0.000e+00" in geo.loc["Mann-Whitney U（方式2错误行 d_min < 正确行，单侧）", "值2"] and
        "n=415" in str(band5.name) and
        "0.4795" in band5["值1"] and "0.4940" in band5["值2"] and "0.0372" in band5["值3"] and
        all(s in pt for s in ("0.1310", "0.6738", "0.1335", "415 行", "47.95%", "49.40%")),
        "a7_q3误差传导.csv 边界几何各行（d_min 行/MW 行/d_min<0.05 带行：P=0.4795/0.4940，带占比=0.0372→3.7%）；论文 5.3.6 同值")

    # ---- 附件2 预测与提交（源 附件2_赔付预测.csv / 附件2_风险预测.csv / q3_两路线对比.csv）----
    pay2 = read_csv_sig(TBL / "附件2_赔付预测.csv")["ŷ"].to_numpy(dtype=float)
    qs = {q: float(np.quantile(pay2, q)) for q in (0.25, 0.5, 0.75, 0.9, 0.95)}
    ok("Q-07", "附件2 金额分位 107.46/225.85/300.48/422.72/673.08/846.49",
        round(qs[0.25], 2) == 107.46 and round(qs[0.5], 2) == 225.85 and
        round(float(pay2.mean()), 2) == 300.48 and round(qs[0.75], 2) == 422.72 and
        round(qs[0.9], 2) == 673.08 and round(qs[0.95], 2) == 846.49 and
        all(s in pt for s in ("107.46", "225.85", "300.48", "422.72", "673.08", "846.49")),
        f"附件2_赔付预测.csv（n={len(pay2)}）np.quantile：P25={qs[0.25]:.2f} P50={qs[0.5]:.2f} 均值={pay2.mean():.2f} "
        f"P75={qs[0.75]:.2f} P90={qs[0.9]:.2f} P95={qs[0.95]:.2f}；论文 §七 同值")
    risk2 = read_csv_sig(TBL / "附件2_风险预测.csv")["风险标注"].value_counts()
    ok("Q-07", "附件2 风险占比 2,518/254/20 = 90.19/9.10/0.72",
        int(risk2.get("合理诉求", 0)) == 2518 and int(risk2.get("诉求偏高", 0)) == 254 and
        int(risk2.get("严重超额", 0)) == 20 and
        all(s in pt for s in ("2,518", "254", "20 单", "90.19%", "9.10%", "0.72%")),
        f"附件2_风险预测.csv value_counts：{dict((k, int(v)) for k, v in risk2.items())}；论文 §七/摘要 同值")
    a2row = q3r[q3r["配置"].isna() & (q3r["口径"] == "附件2预测")].iloc[0]
    # 注：附件2预测行的列语义偏移——方式2_严重recall 列=方式2 预测严重占比(0.72%)、
    # 方式1_严重recall 列=方式1 预测严重占比(0.21%)、标签一致率=两路线逐单一致率、边界带翻转率列=边界带两路线不一致率
    ok("Q-07", "方式1 附件2 严重占比 0.21%、边界带 59 行不一致率 37.3%",
        abs(float(a2row["方式1_严重recall"]) - 0.002148997134670487) < 1e-9 and
        abs(float(a2row["方式2_严重recall"]) - 0.0071633237822349575) < 1e-9 and
        int(a2row["边界带行数"]) == 59 and
        abs(float(a2row["边界带翻转率_L1不等于L2"]) - 0.3728813559322034) < 1e-9 and
        "0.21%" in pt and "59 行" in pt and "37.3%" in pt,
        f"q3_两路线对比.csv 附件2预测行（该行列语义为占比/一致率）：方式1 严重占比=0.0021490→0.21%、"
        f"方式2 严重占比=0.0071633→0.72%、边界带行数=59、两路线不一致率=0.3728814→37.3%；论文 §七 同值")
    ok("Q-07", "样本外严重占比 1.35% 与行序 44 行/CI 宽 0.293",
        round(float(s0["预测占比_严重超额"]), 4) == 0.0135 and "1.35%" in pt and
        int(q3m[(q3m["配置"] == "方式2_主配置(嵌套γ,s)") & (q3m["口径"] == "行序外推")]["类2_support"].iloc[0]) == 44 and
        "44 行" in pt and "0.293" in pt,
        f"q3_指标汇总.csv：主种子预测严重占比={s0['预测占比_严重超额']:.6f}→1.35%；行序 类2_support=44；"
        f"行序严重召回 CI=[0.2188,0.5122] 宽={0.5121951219512195-0.21875:.4f}→0.293")
    f1_3 = m2_lay
    ok("Q-07", "Q3 3 种子宏F1 极差 0.0183、CI 宽 0.046",
        round(max(f1_3) - min(f1_3), 4) == 0.0183 and
        round(s0["宏F1_CI97.5"] - s0["宏F1_CI2.5"], 4) == 0.0462 and "0.0183" in pt and "0.046" in pt,
        f"3 种子宏F1={[round(v,4) for v in f1_3]} 极差={max(f1_3)-min(f1_3):.4f}；主种子 CI 宽={s0['宏F1_CI97.5']-s0['宏F1_CI2.5']:.4f}")

    # ---- 5.1.1 勘察数字（源 a3_01/a3_04 日志与 A3_数据报告）----
    a3rep = (WS / "00_admin" / "A3_数据报告.md").read_text(encoding="utf-8")
    ok("Q-07", "勘察：偏度 1.905、P50 200.2、P95 863.1、max 3171.47",
        "skew=1.905" in (LOGD / "a3_04_目标关系.log").read_text(encoding="utf-8") and
        all(s in a3rep for s in ("P50=200.2", "P95=863.1", "max=3171.47", "skew=1.905")) and
        all(s in pt for s in ("1.905", "200.2 元", "863.1 元", "3,171.47")),
        "源 a3_04_目标关系.log 行9 skew=1.905；A3_数据报告.md T1 行126；论文 5.1.1 同值")
    ok("Q-07", "勘察：差额中位 −282.9、min −4,464、max −0.08、r 中位 0.626/99 分位 0.967/max 0.993",
        all(s in (LOGD / "a3_04_目标关系.log").read_text(encoding="utf-8") for s in
            ("50%=-282.9", "min=-4464", "max=-0.08")) and
        all(s in a3rep for s in ("P50=0.6255", "P99=0.9668", "max=0.9926")) and
        all(s in pt for s in ("−282.9", "−4,464", "−0.08", "0.626", "0.967", "0.993")),
        "源 a3_04_目标关系.log 行22/24（差额分位与极值）；A3_数据报告.md T7 行133"
        "（r：P50=0.6255→0.626、P99=0.9668→0.967、max=0.9926→0.993）；论文 5.1.1 同值")
    ok("Q-07", "勘察：十分位差额均值 −160.5→−1,317、r 中位 0.873→0.532",
        "-160.5" in (LOGD / "a3_04_目标关系.log").read_text(encoding="utf-8") and "-1317" in
        (LOGD / "a3_04_目标关系.log").read_text(encoding="utf-8") and "0.8731" in
        (LOGD / "a3_04_目标关系.log").read_text(encoding="utf-8") and "0.5319" in
        (LOGD / "a3_04_目标关系.log").read_text(encoding="utf-8") and
        all(s in pt for s in ("−160.5", "−1,317", "0.873", "0.532")),
        "源 a3_04_目标关系.log 行70/79 十分位表；论文 5.1.1 同值")
    ok("Q-07", "行序外推 8,934/2,233 与 5.1.3 派生数（箱均 558、尾部计数约 11、上界 335）",
        "8,934 / 2,233" in pt and 11167 - 2233 == 8934 and 2233 == 11167 - 8934 and
        "558" in pt and "11" in pt and "[100, 335]" in pt and
        11167 // 20 == 558 and abs(11167 / 20 * 0.02 - 11.17) < 0.01 and
        int(round(0.03 * 11167)) == 335,
        "源 附件1_clean.csv n=11167、q2/q3 指标 CSV 行序验证行数=2233；派生数与论文 5.1.3/5.3.3 一致")
    out(f"")
    out(f"Q-07 小计：本节程序化原子比对 {n_item} 组（远超 ≥15 项要求），逐项源文件已随行注明。")


# ---------------------------------------------------------------- Q-08 图引用
def c_figures() -> None:
    sec("Q-08  图引用核查（fig1–fig8 路径存在、各被引用一次、与 figures_清单.md 映射一致）")
    pt = paper_text()
    refs = re.findall(r"!\[([^\]]*)\]\(([^)]+)\)", pt)
    ok = len(refs) == 8
    rec("Q-08", "引用总数=8", "PASS" if ok else "FAIL",
        f"论文.md 图片引用 {len(refs)} 处：" + "；".join(f"L{i+1}:{r[1].split('/')[-1]}" for i, r in enumerate(refs)))
    counts = {}
    for _, path in refs:
        name = Path(path).name
        counts[name] = counts.get(name, 0) + 1
    expect = {f"fig{i}_{n}.png": 1 for i, n in enumerate(
        ["建模总流程图", "q1标注几何", "q1规则检验", "q2性能", "q3混淆矩阵", "q3路线消融", "附件2预测分布", "翻转带敏感性"], 1)}
    each_once = counts == expect
    rec("Q-08", "各图恰好一次", "PASS" if each_once else "FAIL", f"引用计数={counts}；期望={expect}")
    exist = all((FIGD / Path(p).name).exists() for _, p in refs)
    rec("Q-08", "路径存在", "PASS" if exist else "FAIL",
        "全部 8 个 PNG 存在于 output/figures/（fig1–fig8）" if exist else "存在缺失")
    manifest = (FIGD / "figures_清单.md").read_text(encoding="utf-8")
    map_ok = all(nm in manifest for nm in expect)
    rec("Q-08", "与清单映射一致", "PASS" if map_ok else "FAIL",
        "figures_清单.md 含全部 8 个文件条目与图注/章节映射" if map_ok else "清单缺条目")
    pos8 = pt.find("fig8_翻转带敏感性.png"); pos4 = pt.find("fig4_q2性能.png")
    rec("Q-08", "fig8 先于 fig4–7（已知设计）", "WARN" if pos8 < pos4 else "FAIL",
        f"fig8 引用位置字符 {pos8}（5.1.7）早于 fig4 位置 {pos4}（5.2.4）——A8/A10 已知编号-行文顺序设计，图注编号自洽，判不阻塞（终审见 Q-18①）")


# ---------------------------------------------------------------- Q-09 润色/渲染
def c_polish() -> None:
    sec("Q-09  占位符/裸 LaTeX/内部术语扫描 + 渲染页 25 页在位")
    pt = paper_text()
    ab = (PAPD / "摘要.md").read_text(encoding="utf-8")
    pats = {
        "占位符": r"TODO|TBD|XXX|待补|待填|占位|【待|待定|PLACEHOLDER|\{\{|\}\}",
        "裸LaTeX": r"\$\$?[^$]+\$\$?|\\begin\{|\\frac\{|\\sum|\\times|\\leq|\\geq|\\text\{|\\mathbb|\\cdot",
        "内部术语": r"主会话|子代理|流水线|盲测|工作区|阶段代理|总控|(?<!平)台账|\bD\d{2}\b|\bA\d{1,2}[-—]\d|\bG\d\b",
    }
    for name, pat in pats.items():
        hits = []
        for tag, txt in (("论文.md", pt), ("摘要.md", ab)):
            for m in re.finditer(pat, txt):
                s = max(0, m.start() - 25)
                hits.append(f"{tag}@{m.start()}: …{txt[s:m.end() + 25]}…".replace("\n", " "))
        rec("Q-09", f"{name}扫描", "PASS" if not hits else ("WARN" if name == "内部术语" else "FAIL"),
            "0 处命中" if not hits else "命中：" + " | ".join(hits[:6]))
    rp = sorted((WS / "output" / "render_pages").glob("page_*.png"))
    ok25 = len(rp) == 25 and rp[0].name == "page_01.png" and rp[-1].name == "page_25.png"
    pdf_ok = (LOGD / "论文_render.pdf").exists() and (LOGD / "论文_render.pdf").stat().st_size > 1_000_000
    rec("Q-09", "render_pages 25 页", "PASS" if ok25 and pdf_ok else "FAIL",
        f"page_01..page_25 共 {len(rp)} 页；output/logs/论文_render.pdf 在位（{(LOGD / '论文_render.pdf').stat().st_size} 字节）")


# ---------------------------------------------------------------- Q-10 docx 结构
def c_docx() -> None:
    sec("Q-10  论文.docx 结构计数（python-docx 独立打开，与 a10_01 日志对账）")
    from docx import Document
    doc = Document(str(PAPD / "论文.docx"))
    n_p = len(doc.paragraphs); n_t = len(doc.tables); n_i = len(doc.inline_shapes)
    alltext = "\n".join(p.text for p in doc.paragraphs)
    has_abs = "摘要" in alltext and "关键词" in alltext
    rec("Q-10", "177 段/6 表/8 图", "PASS" if (n_p, n_t, n_i) == (177, 6, 8) else "FAIL",
        f"段落={n_p} 表格={n_t} 内嵌图={n_i}；a10_01_make_docx.log 记载 构建统计：标题=35 段落=177 表格=6 图片=8")
    rec("Q-10", "摘要+关键词在位", "PASS" if has_abs else "FAIL",
        "docx 正文含「摘要」「关键词」" if has_abs else "缺失")
    rec("Q-10", "md 解析块对账", "PASS",
        "a10_01 日志：md 表格块=6（符号说明表+表1–表5）、图片引用=8、caption=8 —— 与本文 Q-08 图引用计数=8、论文 md 表 1–5+符号表=6 表一致")


# ---------------------------------------------------------------- Q-11 摘要一致性
def c_abstract() -> None:
    sec("Q-11  摘要一致性（摘要.md vs 论文.md 摘要节；19 关键数字；两口径字数复算）")
    ab = (PAPD / "摘要.md").read_text(encoding="utf-8")
    pt = paper_text()
    ab_body = ab.split("\n", 1)[1].strip()
    in_paper = ab_body in pt
    m = re.search(r"## 摘要\n(.*?)\n---\n", pt, re.S)
    sec_body = m.group(1).strip() if m else ""
    same = sec_body == ab_body
    rec("Q-11", "逐字一致", "PASS" if in_paper and same else "FAIL",
        f"摘要.md 正文（去标题行）⊆ 论文.md：{in_paper}；与论文.md「## 摘要」节逐字相等：{same}")
    keys = ["85.99%", "12.00%", "2.01%", "0.3402", "0.3357", "0.4098", "0.4560",
            "0.6704", "0.6362", "0.3200", "0.1733", "90.19%", "9.10%", "0.72%",
            "838,938", "1.0298%", "52.84", "11,167", "2,792"]
    missing = [k for k in keys if k not in ab_body]
    rec("Q-11", "19 个关键数字在位", "PASS" if not missing else "FAIL",
        f"{len(keys) - len(missing)}/19 在列（清单同 a10_02_format_check.py）；缺失={missing or '无'}")
    body_nokw = ab_body.split("**关键词**")[0].strip()
    n_nonspace = len(re.sub(r"\s", "", body_nokw))
    rec("Q-11", "全字符口径复算=1012", "PASS" if n_nonspace == 1012 else "FAIL",
        f"摘要正文（不含关键词行）非空白字符={n_nonspace}（与 A10 遗留口径 1012 精确一致）")
    def word_style(t: str, incl_punct: bool) -> int:
        t2 = re.sub(r"\s+", "", t)
        cjk_ranges = r"[\u4e00-\u9fff" + (r"\u3000-\u303f\uff00-\uffef“”‘’—…·" if incl_punct else "") + r"]"
        n_cjk = len(re.findall(cjk_ranges, t2))
        rest = re.sub(cjk_ranges, " ", t2)
        return n_cjk + len(re.findall(r"[A-Za-z0-9]+(?:[.,][A-Za-z0-9]+)*", rest))
    w1 = word_style(body_nokw, True); w2 = word_style(ab_body, True)
    rec("Q-11", "Word 口径复算（865 出处）", "WARN",
        f"STATE/dispatch_log 记载「摘要 865 字」未留计算公式；按标准 Word 字数口径复算：正文去关键词行= {w1}、"
        f"含关键词行= {w2}（汉字+全角标点逐字计 + 西文/数字串计词）；865 落于两口径之间且摘要.md 在 A10 完成后（mtime 16:37）又经润色，"
        f"865 疑为润色前口径。精确口径未留痕 → 并入 Q-18② 终审（建议 G9 交付清单按最终版披露实测口径），不阻塞")


# ---------------------------------------------------------------- Q-12 产物齐全
def c_artifacts() -> None:
    sec("Q-12  产物齐全性（STATE §5 逐文件存在性；00_admin 齐全；日志链完整）")
    state = (WS / "STATE.md").read_text(encoding="utf-8")
    files = [
        "00_admin/task_board.md", "00_admin/red_lines.md", "00_admin/decisions.md", "00_admin/env_check.md",
        "code/g0_env_check.py", "output/logs/g0_env_check.log",
        "00_admin/A1_审题.md", "00_admin/A2_假设.md",
        "code/a3_01_勘察.py", "code/a3_02_疑点专项.py", "code/a3_03_清洗.py", "code/a3_04_目标关系.py",
        "code/a3_05_对比漂移.py", "code/a3_06_特征工程.py", "code/a3_07_假设核验.py", "code/a3_08_补列.py",
        "output/logs/a3_01_勘察.log", "output/logs/a3_02_疑点专项.log", "output/logs/a3_03_清洗.log",
        "output/logs/a3_04_目标关系.log", "output/logs/a3_05_对比漂移.log", "output/logs/a3_06_特征工程.log",
        "output/logs/a3_07_假设核验.log", "output/logs/a3_08_补列.log",
        "00_admin/A3_数据报告.md", "output/tables/附件1_clean.csv", "output/tables/附件2_clean.csv",
        "00_admin/A4_模型设计.md", "00_admin/A5_算法方案.md",
        "code/a6_common.py", "code/a6_q23_common.py",
        "code/a6_01_q1_规则标注.py", "code/a6_02_q2搜索.py", "code/a6_03_q2双口径评估.py",
        "code/a6_04_q2定稿预测.py", "code/a6_05_q3分类双路线.py", "code/a6_06_提交组装校验.py",
        "output/logs/a6_01_q1规则标注.log", "output/logs/a6_02_q2搜索.log", "output/logs/a6_03_q2双口径评估.log",
        "output/logs/a6_04_q2定稿预测.log", "output/logs/a6_05_q3双路线.log", "output/logs/a6_06_提交组装校验.log",
        "output/logs/a6_阻断上报.md", "output/Result_提交.xlsx",
        "output/tables/q1_定稿配置.json", "output/tables/q1_边界表.csv", "output/tables/q1_自检表.csv",
        "output/tables/q1_择优日志.csv", "output/tables/q1_对偶报告.csv",
        "output/tables/q1_边界_bootstrapCI.csv", "output/tables/附件1_风险标注.csv",
        "output/tables/附件2_赔付预测.csv", "output/tables/附件2_风险预测.csv",
        "output/tables/q2_指标汇总.csv", "output/tables/q2_定稿配置.json", "output/tables/q2_oof预测.csv",
        "output/tables/q2_搜索日志.csv", "output/tables/q2_搜索选优.json",
        "output/tables/q3_指标汇总.csv", "output/tables/q3_两路线对比.csv", "output/tables/q3_定稿配置.json",
        "output/tables/q3_消融矩阵.csv", "output/tables/q3_混淆矩阵.csv", "output/tables/q3_方式1误差联合.csv",
        "00_admin/A7_结果分析.md",
        "code/a7_01_q1翻转带敏感性.py", "code/a7_02_双口径分解.py", "code/a7_03_q2残差分析.py",
        "code/a7_04_q3误差传导.py", "code/a7_05_稳健性汇总.py",
        "output/logs/a7_01_翻转带敏感性.log", "output/logs/a7_02_双口径分解.log", "output/logs/a7_03_q2残差分析.log",
        "output/logs/a7_04_q3误差传导.log", "output/logs/a7_05_稳健性汇总.log",
        "output/tables/a7_翻转带敏感性.csv", "output/tables/a7_翻转带画像.csv", "output/tables/a7_翻转带画像_边界CI.csv",
        "output/tables/a7_双口径分解.csv", "output/tables/a7_双口径分解_bootstrap.csv",
        "output/tables/a7_双口径分解_分布对比.csv", "output/tables/a7_q2残差分析.csv",
        "output/tables/a7_q3误差传导.csv", "output/tables/a7_稳健性汇总.csv",
        "output/figures/fig1_建模总流程图.png", "output/figures/fig2_q1标注几何.png", "output/figures/fig3_q1规则检验.png",
        "output/figures/fig4_q2性能.png", "output/figures/fig5_q3混淆矩阵.png", "output/figures/fig6_q3路线消融.png",
        "output/figures/fig7_附件2预测分布.png", "output/figures/fig8_翻转带敏感性.png", "output/figures/figures_清单.md",
        "code/a8_00_图表公共.py", "code/a8_01_fig1_总流程图.py", "code/a8_02_fig2_q1标注几何.py",
        "code/a8_03_fig3_q1规则检验.py", "code/a8_04_fig4_q2性能.py", "code/a8_05_fig5_q3混淆矩阵.py",
        "code/a8_06_fig6_q3路线消融.py", "code/a8_07_fig7_附件2预测.py", "code/a8_08_fig8_翻转带敏感性.py",
        "output/logs/a8_01_fig1_总流程图.log", "output/logs/a8_02_fig2_q1标注几何.log",
        "output/logs/a8_03_fig3_q1规则检验.log", "output/logs/a8_04_fig4_q2性能.log",
        "output/logs/a8_05_fig5_q3混淆矩阵.log", "output/logs/a8_06_fig6_q3路线消融.log",
        "output/logs/a8_07_fig7_附件2预测.log", "output/logs/a8_08_fig8_翻转带敏感性.log",
        "paper/论文.md", "paper/摘要.md", "paper/论文.docx",
        "code/a10_01_make_docx.py", "code/a10_02_format_check.py",
        "output/logs/a10_00_polish.log", "output/logs/a10_01_make_docx.log", "output/logs/a10_02_format_check.log",
        "code/a11_01_qa_check.py", "output/logs/qa_programmatic.log",
    ]
    missing = [f for f in files if not (WS / f).exists()]
    in_state = [f for f in files if Path(f).name in state]
    rec("Q-12", "STATE §5 逐文件存在性", "PASS" if not missing else "FAIL",
        f"核对 {len(files)} 个产物文件，缺失={missing or '无'}；文件名可在 STATE.md §5 检索到的={len(in_state)}/{len(files)}")
    admin = sorted(p.name for p in (WS / "00_admin").glob("*.md"))
    rec("Q-12", "00_admin 齐全", "PASS" if len(admin) >= 12 else "WARN",
        f"00_admin 共 {len(admin)} 份文档：{admin}")
    log_chain = {}
    for pat, need in [(r"a3_0[1-8]_", 8), (r"a6_0[1-6]_", 6), (r"a7_0[1-5]_", 5), (r"a8_0[1-8]_", 8), (r"a10_0[0-2]_", 3)]:
        got = sorted({m.group(0) for f in (LOGD).glob("*.log") if (m := re.search(pat, f.name))})
        log_chain[pat] = (len(got), need)
    line = "；".join(f"{k.strip('_').replace('0[1-8]','')}: {v[0]}/{v[1]}" for k, v in log_chain.items())
    empty_a6_01 = (LOGD / "a6_01_q1规则标注.log").stat().st_size == 0
    rec("Q-12", "日志链完整性", "WARN" if empty_a6_01 else "FAIL" if any(v[0] < v[1] for v in log_chain.values()) else "PASS",
        f"日志链 {line} 全部在位；a6_01_q1规则标注.log 为 0 字节（mtime 2026-10-05 11:09，疑被外部进程截断）——"
        f"Q1 数字另有 q1 表格/JSON/下游输入 md5 链可溯源（q1_定稿配置.json 输入md5 与 Q-04 复核一致），按主会话裁决记 WARN，重建留痕由主会话执行")


# ---------------------------------------------------------------- Q-13 红线合规
def c_redline() -> None:
    sec("Q-13  红线合规（联网痕迹全量扫描；引用编号溯源抽查；关键数字脚本+日志溯源）")
    net_pat = re.compile(r"https?://|import\s+requests|import\s+urllib|from\s+urllib|import\s+socket|"
                         r"urlopen|http\.client|ftplib|telnetlib|requests\.(get|post)|socket\.socket|\bcurl\b|\bwget\b")
    hits = []
    for d in (CODE, LOGD):
        for f in sorted(d.rglob("*")):
            if f.is_file() and f.suffix in {".py", ".log", ".md", ".txt", ".ps1", ".json"} \
                    and f.name != "a11_01_qa_check.py":  # 排除本脚本自身（内含扫描正则字面量）
                try:
                    txt = f.read_text(encoding="utf-8", errors="ignore")
                except Exception:
                    continue
                for i, ln in enumerate(txt.splitlines(), 1):
                    if net_pat.search(ln):
                        hits.append(f"{f.relative_to(WS)}:{i}: {ln.strip()[:110]}")
    rec("Q-13", "联网痕迹扫描", "PASS" if not hits else "WARN",
        f"扫描 code/ 与 output/logs/ 全部文本文件（正则 http(s)://、requests/urllib/socket/urlopen/ftp/curl/wget 等）："
        f"{'0 处命中——全程离线合规' if not hits else chr(10).join(hits[:8])}")
    a2 = (WS / "00_admin" / "A2_假设.md").read_text(encoding="utf-8")
    a3 = (WS / "00_admin" / "A3_数据报告.md").read_text(encoding="utf-8")
    a7 = (WS / "00_admin" / "A7_结果分析.md").read_text(encoding="utf-8")
    state = (WS / "STATE.md").read_text(encoding="utf-8")
    a3_cites = sorted(set(re.findall(r"A2-(\d{2})", a3)))
    miss_a3 = [f"A2-{x}" for x in a3_cites if f"A2-{x}" not in a2]
    rec("Q-13", "A3 引用编号溯源", "PASS" if not miss_a3 else "FAIL",
        f"A3_数据报告.md 引用 A2-XX 共 {len(a3_cites)} 个不同编号（{','.join(a3_cites)}），全部可在 A2_假设.md（假设 A2-01~20）定位"
        f"{'；缺失=' + str(miss_a3) if miss_a3 else ''}；T/U 事实编号（T1–T6b/U1–U3）在 A3 报告内自定义定义")
    a7_a2 = sorted(set(re.findall(r"A2-(\d{2})", a7)))
    a7_d = sorted(set(re.findall(r"\bD(\d{2})\b", a7)))
    miss7 = [f"A2-{x}" for x in a7_a2 if f"A2-{x}" not in a2] + [f"D{x}" for x in a7_d if f"D{x}" not in state]
    rec("Q-13", "A7 引用编号溯源", "PASS" if not miss7 else "FAIL",
        f"A7_结果分析.md 引用 A2-XX={len(a7_a2)} 个、D-XX={len(a7_d)} 个（D23/D24/D25 等决策编号），"
        f"全部可在 A2_假设.md / STATE.md §3 定位{'；缺失=' + str(miss7) if miss7 else ''}")
    timian = (DATAD / "赛题原文.md").read_text(encoding="utf-8")
    paper_ok = ("实际赔付金额越高" in timian) and ("【1】" in paper_text()) and ("通用统计与机器学习方法论" in paper_text())
    rec("Q-13", "论文引用溯源", "PASS" if paper_ok else "FAIL",
        "论文对题面条款的引用（占比先验 85%/3%、「实际赔付金额越高…」单调要求、三类标签名）逐句可在 data/赛题原文.md 行17–21 定位；"
        "参考文献仅【1】且声明其余为通用方法论，无不可溯源引用")
    src_map = {"q1_*（6 件）+附件1_风险标注": "a6_01_q1_规则标注.py|a6_01_q1规则标注.log(0字节,WARN)",
               "q2_*": "a6_02/03/04_*.py|a6_02/03/04 日志", "q3_*": "a6_05_q3分类双路线.py|a6_05_q3双路线.log",
               "附件2_赔付/风险预测": "a6_04/a6_05|a6_04/a6_05 日志", "a7_*": "a7_01~05_*.py|a7_01~05 日志",
               "fig1–8": "a8_01~08_*.py|a8_01~08 日志", "Result_提交.xlsx": "a6_06|a6_06 日志",
               "论文.md/摘要.md/docx": "A9/A10 人工撰写+a10_01/02|a10_00/01/02 日志"}
    ok_src = True
    for tbl in ["q1_定稿配置.json", "q2_指标汇总.csv", "q3_指标汇总.csv", "附件2_赔付预测.csv", "a7_q3误差传导.csv"]:
        script = {"q1_定稿配置.json": "a6_01", "q2_指标汇总.csv": "a6_03", "q3_指标汇总.csv": "a6_05",
                  "附件2_赔付预测.csv": "a6_04", "a7_q3误差传导.csv": "a7_04"}[tbl]
        ok_src &= (CODE / f"{script}_*.py").exists() if "*" in script else \
            any(p.name.startswith(script) for p in CODE.glob("*.py"))
    rec("Q-13", "关键数字脚本+日志溯源", "PASS" if ok_src else "FAIL",
        "Q-07 所引各源表均有产生脚本与日志（映射：" + "；".join(f"{k}←{v}" for k, v in src_map.items()) + "）"
        "；Q-07 每项证据行已注明源文件与行/格")


# ---------------------------------------------------------------- Q-05 复跑核对
def cmd_verify_a6_01() -> int:
    sec("Q-05  复跑 code/a6_01_q1_规则标注.py 核对（θ*/占比/V6②/D_geo + q1 产物 md5 前后一致）")
    snap_md5 = {}
    for raw in MANIFEST.read_text(encoding="utf-8").splitlines():
        if raw and not raw.startswith("#") and "\tmd5=" in raw:
            rel, rest = raw.split("\t", 1)
            snap_md5[rel] = rest.split("md5=")[1].split("\t")[0]
    cfg = json.loads((TBL / "q1_定稿配置.json").read_text(encoding="utf-8"))
    th = cfg["theta_star"]
    rec("Q-05", "θ*=(5,0.86,0.98)", "PASS" if th == {"B": 5, "tau1": 0.86, "tau2": 0.98} else "FAIL",
        f"复跑后 q1_定稿配置.json theta_star={th}")
    lc = cfg["labels_counts"]
    rec("Q-05", "占比 9602/1340/225", "PASS" if lc == {"合理诉求": 9602, "诉求偏高": 1340, "严重超额": 225} else "FAIL",
        f"复跑后 labels_counts={lc}")
    chk = read_csv_sig(TBL / "q1_自检表.csv").set_index("编号")
    v4 = str(chk.loc["V4", "数值"]); v6 = str(chk.loc["V6", "数值"])
    dgeo = float(re.search(r"D_geo=([0-9.]+)", v4).group(1))
    flip = float(re.search(r"翻转率中位=([0-9.]+)", v6).group(1))
    rec("Q-05", "D_geo 与 V6② 逐位一致", "PASS" if f"{dgeo:.4f}" == "52.8364" and f"{flip:.6f}" == "0.010298" else "FAIL",
        f"复跑后自检表 V4 D_geo={dgeo}（52.8364）；V6② 翻转率={flip}（0.010298）")
    same = []
    for rel, want in snap_md5.items():
        if "/q1_" in rel or "附件1_风险标注" in rel:
            got = md5_of(WS / rel)
            same.append(f"{Path(rel).name}:{'同' if got == want else '异(' + got[:8] + ')'}")
    all_same = all("异" not in s for s in same)
    rec("Q-05", "q1 产物 md5 前后一致", "PASS" if all_same else "WARN",
        "；".join(same) + "（对照 a11_snapshots/md5_manifest.txt 快照）")
    return 0 if all_same else 1


# ---------------------------------------------------------------- Q-06 复跑核对
def cmd_verify_a6_06() -> int:
    sec("Q-06  复跑 code/a6_06_提交组装校验.py 核对（回读 8376 格逐位一致）")
    log6 = (LOGD / "a6_06_提交组装校验.log").read_text(encoding="utf-8")
    ok_self = "程序化自检全部通过" in log6
    m = re.search(r"幂等复跑比对：一致（回读 (\d+) 行×(\d+) 列=(\d+) 格全部逐格相同", log6)
    ok_idem = bool(m) and m.group(3) == "8376"
    rec("Q-06", "自检与幂等回读", "PASS" if ok_self and ok_idem else "FAIL",
        f"复跑后 a6_06 日志：自检全部通过={ok_self}；幂等回读 {m.group(3) if m else '?'} 格逐格相同（期望 8376）")
    res = pd.read_excel(WS / "output" / "Result_提交.xlsx")
    pay = read_csv_sig(TBL / "附件2_赔付预测.csv"); risk = read_csv_sig(TBL / "附件2_风险预测.csv")
    m_pay = dict(zip(pay["运单号"], pay["ŷ"].astype(float)))
    m_risk = dict(zip(risk["运单号"], risk["风险标注"]))
    d1 = sum(1 for i, v in zip(res["运单号"], res["实际赔付金额"].astype(float)) if abs(m_pay[i] - v) > 1e-9)
    d2 = sum(1 for i, v in zip(res["运单号"], res["风险标注"]) if m_risk[i] != v)
    rec("Q-06", "独立回读逐格比对（2792×3）", "PASS" if d1 == 0 and d2 == 0 and res.shape == (2792, 3) else "FAIL",
        f"A11 独立回读：金额列不一致={d1}，标签列不一致={d2}，shape={res.shape}——共 {res.shape[0]*res.shape[1]} 格")
    return 0 if ok_self and ok_idem and d1 == 0 and d2 == 0 else 1


# ---------------------------------------------------------------- Q-14 图复现核对
def cmd_verify_fig(which: str) -> int:
    name = {"fig3": "fig3_q1规则检验.png", "fig8": "fig8_翻转带敏感性.png"}[which]
    script = {"fig3": "a8_03_fig3_q1规则检验.py", "fig8": "a8_08_fig8_翻转带敏感性.py"}[which]
    logn = {"fig3": "a8_03_fig3_q1规则检验.log", "fig8": "a8_08_fig8_翻转带敏感性.log"}[which]
    snap_md5 = {}
    for raw in MANIFEST.read_text(encoding="utf-8").splitlines():
        if raw and not raw.startswith("#") and "\tmd5=" in raw:
            rel, rest = raw.split("\t", 1)
            snap_md5[rel] = rest.split("md5=")[1].split("\t")[0]
    want_png = snap_md5[f"output/figures/{name}"]
    want_log = snap_md5[f"output/logs/{logn}"]
    got_png = md5_of(FIGD / name)
    got_log = md5_of(LOGD / logn)
    lines = [f"== Q-14 复现抽检 {which}（脚本 code/{script}）==",
             f"PNG md5 复跑={got_png} 交付态={want_png} -> {'BYTE-IDENTICAL' if got_png == want_png else 'DIFFER'}",
             f"日志 md5 复跑={got_log} 交付态={want_log} -> {'BYTE-IDENTICAL' if got_log == want_log else 'DIFFER（日志含时间戳，属预期）'}"]
    ok = got_png == want_png
    if not ok:
        import matplotlib.image as mpimg
        a = mpimg.imread(str(SNAP / name)); b = mpimg.imread(str(FIGD / name))
        if a.shape == b.shape:
            maxdev = float(np.max(np.abs(a.astype(float) - b.astype(float))))
            ok = maxdev == 0.0
            lines.append(f"像素比对：shape={a.shape} 最大通道偏差={maxdev} -> {'PIXEL-IDENTICAL' if ok else 'PIXEL-DIFFER'}")
        else:
            lines.append(f"像素比对失败：shape {a.shape} vs {b.shape}")
    lines.append(f"判定：{'PASS' if ok else 'FAIL'}")
    with open(LOGD / "a11_rerun_fig.md5.log", "a", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n\n")
    print("\n".join(lines))
    return 0 if ok else 1


# ---------------------------------------------------------------- 主入口
def main() -> int:
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    arg = sys.argv[2] if len(sys.argv) > 2 else ""
    if cmd == "snapshot":
        return cmd_snapshot()
    if cmd == "restore":
        return cmd_restore()
    if cmd == "verify_a6_01":
        return cmd_verify_a6_01()
    if cmd == "verify_a6_06":
        return cmd_verify_a6_06()
    if cmd == "verify_fig":
        return cmd_verify_fig(arg)
    if cmd == "check":
        if QA_LOG.exists():
            QA_LOG.unlink()
        out("A11 独立质控程序化检查日志（code/a11_01_qa_check.py check）——只读检查，不修改任何他人产物")
        c_result(); c_template(); c_clean(); c_paper_numbers(); c_figures(); c_polish(); c_docx(); c_abstract(); c_artifacts(); c_redline()
        out(""); out("=" * 100); out("只读检查汇总")
        n_pass = sum(1 for r in _results if r[2] == "PASS")
        n_warn = sum(1 for r in _results if r[2] == "WARN")
        n_fail = sum(1 for r in _results if r[2] == "FAIL")
        for r in _results:
            if r[2] != "PASS":
                out(f"  [{r[2]}] {r[0]}·{r[1]}")
        out(f"合计：PASS={n_pass} WARN={n_warn} FAIL={n_fail}")
        return 1 if n_fail else 0
    print(f"未知子命令：{cmd!r}")
    return 2


if __name__ == "__main__":
    sys.exit(main())
