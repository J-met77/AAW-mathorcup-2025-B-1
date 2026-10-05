# -*- coding: utf-8 -*-
"""
g9_check_清单核对.py —— G9 交付打包与最终清单核对脚本（A0 收尾专用）
=====================================================================
职责：对盲测 Run-2 工作区全部交付物与支撑材料逐文件核对【存在性 + 字节大小 + md5】，
      并对竞赛提交物做轻量结构复验（只读，不修改任何他人产物）。
产出：output/logs/g9_check_清单核对.log（运行日志，逐文件留痕）+ stdout 汇总。
红线：全程离线；只在 agent_workspace_B2 工作区内读写；不修改 STATE.md 与他人产物；
      本脚本为 G9 新增文件之一（另两件：output/交付清单.md 与本日志）。
日期：2026-10-05
"""
import sys
import json
import hashlib
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent          # agent_workspace_B2/
LOG  = BASE / "output" / "logs" / "g9_check_清单核对.log"

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

SELF_LOG_NAME = LOG.name
SELF_NAME = Path(__file__).name

# ---------------------------------------------------------------- 期望清单
# 1) 竞赛提交物（4）
SUBMISSIONS = [
    "output/Result_提交.xlsx",
    "paper/论文.docx",
    "paper/论文.md",
    "paper/摘要.md",
]

# 2) 代码（34 项既有 + 本脚本 = 35）
CODE = [
    "code/g0_env_check.py",
    "code/a3_01_勘察.py", "code/a3_02_疑点专项.py", "code/a3_03_清洗.py",
    "code/a3_04_目标关系.py", "code/a3_05_对比漂移.py", "code/a3_06_特征工程.py",
    "code/a3_07_假设核验.py", "code/a3_08_补列.py",
    "code/a6_common.py", "code/a6_q23_common.py",
    "code/a6_01_q1_规则标注.py", "code/a6_02_q2搜索.py", "code/a6_03_q2双口径评估.py",
    "code/a6_04_q2定稿预测.py", "code/a6_05_q3分类双路线.py", "code/a6_06_提交组装校验.py",
    "code/a7_01_q1翻转带敏感性.py", "code/a7_02_双口径分解.py", "code/a7_03_q2残差分析.py",
    "code/a7_04_q3误差传导.py", "code/a7_05_稳健性汇总.py",
    "code/a8_00_图表公共.py", "code/a8_01_fig1_总流程图.py", "code/a8_02_fig2_q1标注几何.py",
    "code/a8_03_fig3_q1规则检验.py", "code/a8_04_fig4_q2性能.py", "code/a8_05_fig5_q3混淆矩阵.py",
    "code/a8_06_fig6_q3路线消融.py", "code/a8_07_fig7_附件2预测.py", "code/a8_08_fig8_翻转带敏感性.py",
    "code/a10_01_make_docx.py", "code/a10_02_format_check.py",
    "code/a11_01_qa_check.py",
    "code/" + SELF_NAME,
]

# 3) 日志（42 项；本脚本日志另行登记）
LOGS = [
    "output/logs/g0_env_check.log",
    "output/logs/a3_01_勘察.log", "output/logs/a3_02_疑点专项.log", "output/logs/a3_03_清洗.log",
    "output/logs/a3_04_目标关系.log", "output/logs/a3_05_对比漂移.log", "output/logs/a3_06_特征工程.log",
    "output/logs/a3_07_假设核验.log", "output/logs/a3_08_补列.log",
    "output/logs/a6_01_q1规则标注.log", "output/logs/a6_02_q2搜索.log", "output/logs/a6_03_q2双口径评估.log",
    "output/logs/a6_04_q2定稿预测.log", "output/logs/a6_05_q3双路线.log", "output/logs/a6_06_提交组装校验.log",
    "output/logs/a6_阻断上报.md",
    "output/logs/a7_01_翻转带敏感性.log", "output/logs/a7_02_双口径分解.log",
    "output/logs/a7_03_q2残差分析.log", "output/logs/a7_04_q3误差传导.log", "output/logs/a7_05_稳健性汇总.log",
    "output/logs/a8_01_fig1_总流程图.log", "output/logs/a8_02_fig2_q1标注几何.log",
    "output/logs/a8_03_fig3_q1规则检验.log", "output/logs/a8_04_fig4_q2性能.log",
    "output/logs/a8_05_fig5_q3混淆矩阵.log", "output/logs/a8_06_fig6_q3路线消融.log",
    "output/logs/a8_07_fig7_附件2预测.log", "output/logs/a8_08_fig8_翻转带敏感性.log",
    "output/logs/a8_fig4_刷新前.md5",
    "output/logs/a10_00_polish.log", "output/logs/a10_01_make_docx.log", "output/logs/a10_02_format_check.log",
    "output/logs/a11_rerun_a6_01.log", "output/logs/a11_rerun_a6_06.log", "output/logs/a11_rerun_fig.md5.log",
    "output/logs/a11_snapshots/md5_manifest.txt", "output/logs/a11_snapshots/restore_md5_evidence.txt",
    "output/logs/qa_check.md", "output/logs/qa_programmatic.log",
    "output/logs/render_gate.ps1", "output/logs/论文_render.pdf",
]

# 4) 表格（31 项 + sensitivity_a7 存档 9 项）
TABLES = [
    "output/tables/q1_定稿配置.json", "output/tables/q1_对偶报告.csv", "output/tables/q1_择优日志.csv",
    "output/tables/q1_自检表.csv", "output/tables/q1_边界_bootstrapCI.csv", "output/tables/q1_边界表.csv",
    "output/tables/q2_oof预测.csv", "output/tables/q2_定稿配置.json", "output/tables/q2_指标汇总.csv",
    "output/tables/q2_搜索日志.csv", "output/tables/q2_搜索选优.json",
    "output/tables/q3_两路线对比.csv", "output/tables/q3_定稿配置.json", "output/tables/q3_指标汇总.csv",
    "output/tables/q3_方式1误差联合.csv", "output/tables/q3_消融矩阵.csv", "output/tables/q3_混淆矩阵.csv",
    "output/tables/a7_q2残差分析.csv", "output/tables/a7_q3误差传导.csv", "output/tables/a7_双口径分解.csv",
    "output/tables/a7_双口径分解_bootstrap.csv", "output/tables/a7_双口径分解_分布对比.csv",
    "output/tables/a7_稳健性汇总.csv", "output/tables/a7_翻转带敏感性.csv", "output/tables/a7_翻转带画像.csv",
    "output/tables/a7_翻转带画像_边界CI.csv",
    "output/tables/附件1_clean.csv", "output/tables/附件2_clean.csv",
    "output/tables/附件1_风险标注.csv", "output/tables/附件2_赔付预测.csv", "output/tables/附件2_风险预测.csv",
    "output/tables/sensitivity_a7/a7_q2残差分析.csv", "output/tables/sensitivity_a7/a7_q3误差传导.csv",
    "output/tables/sensitivity_a7/a7_双口径分解.csv", "output/tables/sensitivity_a7/a7_双口径分解_bootstrap.csv",
    "output/tables/sensitivity_a7/a7_双口径分解_分布对比.csv", "output/tables/sensitivity_a7/a7_稳健性汇总.csv",
    "output/tables/sensitivity_a7/a7_翻转带敏感性.csv", "output/tables/sensitivity_a7/a7_翻转带画像.csv",
    "output/tables/sensitivity_a7/a7_翻转带画像_边界CI.csv",
]

# 5) 图表（9）
FIGURES = [
    "output/figures/fig1_建模总流程图.png", "output/figures/fig2_q1标注几何.png",
    "output/figures/fig3_q1规则检验.png", "output/figures/fig4_q2性能.png",
    "output/figures/fig5_q3混淆矩阵.png", "output/figures/fig6_q3路线消融.png",
    "output/figures/fig7_附件2预测分布.png", "output/figures/fig8_翻转带敏感性.png",
    "output/figures/figures_清单.md",
]

# 6) 管理文档（12）+ 台账（1）
ADMIN = [
    "00_admin/task_board.md", "00_admin/red_lines.md", "00_admin/decisions.md",
    "00_admin/env_check.md", "00_admin/dispatch_log.md", "00_admin/automation_备份.md",
    "00_admin/A1_审题.md", "00_admin/A2_假设.md", "00_admin/A3_数据报告.md",
    "00_admin/A4_模型设计.md", "00_admin/A5_算法方案.md", "00_admin/A7_结果分析.md",
    "STATE.md",
]

# 7) 论文渲染件（25 页）
RENDER_PAGES = [f"output/render_pages/page_{i:02d}.png" for i in range(1, 26)]

# 8) 输入数据（4，支撑复跑）
DATA = ["data/赛题原文.md", "data/附件1.xlsx", "data/附件2.xlsx", "data/Result.xlsx"]

GROUPS = [
    ("竞赛提交物", SUBMISSIONS), ("代码脚本", CODE), ("运行日志", LOGS),
    ("结果表格", TABLES), ("论文图表", FIGURES), ("管理文档与台账", ADMIN),
    ("论文渲染页", RENDER_PAGES), ("输入数据", DATA),
]

RESULT_MD5_EXPECTED = "a66012eaa89d01b6797f704f3f44d589"   # 源：qa_check.md Q-01 / STATE §4

# ---------------------------------------------------------------- 工具
def md5_of(p: Path, buf=1 << 20) -> str:
    h = hashlib.md5()
    with open(p, "rb") as f:
        while True:
            b = f.read(buf)
            if not b:
                break
            h.update(b)
    return h.hexdigest()

log_lines = []
def emit(s: str = ""):
    print(s)
    log_lines.append(s)

# ---------------------------------------------------------------- 逐文件核对
emit("# G9 逐文件核对日志（code/g9_check_清单核对.py）")
emit("# 工作区：%s" % BASE)
emit("# 核对内容：存在性 + 字节大小 + md5；竞赛提交物另做只读结构复验")
emit("# 红线声明：全程离线；仅读写本工作区；未修改任何他人产物")
emit("")

grand_total = grand_ok = 0
missing, empty, md5_map = [], [], {}
for gname, rels in GROUPS:
    emit("## 组：%s（期望 %d 项）" % (gname, len(rels)))
    g_ok = 0
    for rel in rels:
        grand_total += 1
        p = BASE / rel
        if not p.exists():
            emit("  [MISSING] %-58s" % rel)
            missing.append(rel)
            continue
        size = p.stat().st_size
        digest = md5_of(p)
        md5_map[rel] = digest
        if size == 0:
            emit("  [EMPTY   ] %-58s size=0 md5=%s" % (rel, digest))
            empty.append(rel)
            continue
        grand_ok += 1
        g_ok += 1
        emit("  [OK      ] %-58s size=%9d md5=%s" % (rel, size, digest))
    emit("  -> 组内 OK %d / %d" % (g_ok, len(rels)))
    emit("")

# 自身日志登记（核对结束时已定稿的大小在收尾补记）
emit("## 自身日志（本脚本产出）")
emit("  [SELFLOG ] output/logs/%s（运行结束时统计）" % SELF_LOG_NAME)
emit("")

# ---------------------------------------------------------------- 结构复验（只读）
emit("## 结构复验（竞赛提交物，只读）")
struct_fail = []
try:
    import pandas as pd
    import openpyxl

    # --- Result_提交.xlsx ---
    rpath = BASE / "output/Result_提交.xlsx"
    wb = openpyxl.load_workbook(rpath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header, body = list(rows[0]), rows[1:]
    checks = []
    checks.append(("shape=(2792,3)", len(body) == 2792 and len(header) == 3))
    checks.append(("列名=[运单号,实际赔付金额,风险标注]",
                   header == ["运单号", "实际赔付金额", "风险标注"]))
    wbids = [r[0] for r in body]
    amounts = [float(r[1]) for r in body]
    labels = [r[2] for r in body]
    from collections import Counter
    cnt = Counter(labels)
    checks.append(("标签域={合理诉求,诉求偏高,严重超额}",
                   set(cnt) == {"合理诉求", "诉求偏高", "严重超额"}))
    checks.append(("占比 2518/254/20 = 90.19/9.10/0.72%",
                   cnt["合理诉求"] == 2518 and cnt["诉求偏高"] == 254 and cnt["严重超额"] == 20))
    checks.append(("金额非负两位小数 min=8.12 max=1455.73",
                   min(amounts) >= 0 and min(amounts) == 8.12 and max(amounts) == 1455.73
                   and all(round(a, 2) == a for a in amounts)))
    df2 = pd.read_csv(BASE / "output/tables/附件2_clean.csv", encoding="utf-8-sig")
    checks.append(("运单号与附件2_clean 逐行零改动",
                   list(df2.iloc[:, 0].astype(str)) == [str(w) for w in wbids]))
    tpl = openpyxl.load_workbook(BASE / "data/Result.xlsx", read_only=True, data_only=True)
    trows = list(tpl.active.iter_rows(values_only=True))
    checks.append(("运单号与模板 Result.xlsx 逐行零改动",
                   [str(r[0]) for r in trows[1:]] == [str(w) for w in wbids]))
    total_amt = round(sum(amounts), 2)
    emit("  Result_提交.xlsx 复验：%s；总额=%.2f 元" %
         ("；".join("%s=%s" % (n, "PASS" if ok else "FAIL") for n, ok in checks), total_amt))
    for n, ok in checks:
        if not ok:
            struct_fail.append("Result: " + n)
    md5_res = md5_map["output/Result_提交.xlsx"]
    emit("  Result md5=%s 与 qa_check Q-01 记载 %s -> %s"
         % (md5_res, RESULT_MD5_EXPECTED, "MATCH" if md5_res == RESULT_MD5_EXPECTED else "DIFF"))
    if md5_res != RESULT_MD5_EXPECTED:
        struct_fail.append("Result md5 与 A11 记载不一致")

    # --- 论文.docx ---
    import docx
    d = docx.Document(str(BASE / "paper/论文.docx"))
    n_para, n_tab, n_img = len(d.paragraphs), len(d.tables), len(d.inline_shapes)
    texts = [p.text for p in d.paragraphs]
    has_abs = any("摘要" in t for t in texts)
    has_kw = any("关键词" in t for t in texts)
    emit("  论文.docx 复验：段落=%d 表格=%d 内嵌图=%d 摘要=%s 关键词=%s"
         % (n_para, n_tab, n_img, has_abs, has_kw))
    if n_tab != 6 or n_img != 8:
        struct_fail.append("论文.docx 表/图数与 6 表 8 图不符")
    if not (has_abs and has_kw):
        struct_fail.append("论文.docx 缺摘要/关键词")

    # --- 摘要.md 与 论文.md「## 摘要」节一致性（Q-11 口径） ---
    abs_txt = (BASE / "paper/摘要.md").read_text(encoding="utf-8")
    pap_txt = (BASE / "paper/论文.md").read_text(encoding="utf-8")
    import re
    m = re.search(r"## 摘要\s*\n(.*?)(?=\n## |\n---\s*\n)", pap_txt, re.S)
    abs_body = "\n".join(abs_txt.splitlines()[1:]) if abs_txt.splitlines() else ""
    norm = lambda s: re.sub(r"\s+", "", s or "")
    eq = bool(m) and norm(abs_body) == norm(m.group(1))
    wc_nonws = len(norm(abs_body))
    emit("  摘要一致性：摘要.md 正文 与 论文.md 摘要节 逐字相等=%s；含关键词行非空白字符=%d"
         % (eq, wc_nonws))
    emit("  （对齐 A11 Q-18② 口径：正文去关键词行 1012、含关键词行 1049；差异仅允许为章节分隔符）")
    if not eq:
        struct_fail.append("摘要.md 与 论文.md 摘要节不一致")

    # --- 论文.md 图表引用计数 ---
    n_figref = len(re.findall(r"output/figures/fig\d_", pap_txt))
    emit("  论文.md 图片引用=%d 处（期望 8，Q-08 口径）" % n_figref)
    if n_figref != 8:
        struct_fail.append("论文.md 图片引用数 != 8")

    # --- 关键数字交叉核对（与源表逐位） ---
    q2 = pd.read_csv(BASE / "output/tables/q2_指标汇总.csv", encoding="utf-8-sig")
    va = q2[(q2["模型"] == "LGBM_变体A")]
    w_strat = va[va["口径"] == "分层5折"]["WAPE"].mean()
    w_seq = va[va["口径"] == "行序外推"]["WAPE"].iloc[0]
    ok_q2 = (round(w_strat, 4) == 0.3402) and (round(w_seq, 4) == 0.3357)
    emit("  Q2 交叉核对：变体A 分层3种子均值 WAPE=%.4f（期望 0.3402）、行序=%.4f（期望 0.3357）-> %s"
         % (w_strat, w_seq, "PASS" if ok_q2 else "FAIL"))
    if not ok_q2:
        struct_fail.append("Q2 WAPE 交叉核对不符")

    q3 = pd.read_csv(BASE / "output/tables/q3_指标汇总.csv", encoding="utf-8-sig")
    f1_m2 = q3[(q3["配置"] == "方式2_主配置(嵌套γ,s)") & (q3["口径"] == "分层5折")]["宏F1"].iloc[0]
    f1_seq = q3[(q3["配置"] == "方式2_主配置(嵌套γ,s)") & (q3["口径"] == "行序外推")]["宏F1"].iloc[0]
    ok_q3 = (round(f1_m2, 4) == 0.6704) and (round(f1_seq, 4) == 0.6710)
    emit("  Q3 交叉核对：方式2 分层主种子 宏F1=%.4f（期望 0.6704）、行序=%.4f（期望 0.6710）-> %s"
         % (f1_m2, f1_seq, "PASS" if ok_q3 else "FAIL"))
    if not ok_q3:
        struct_fail.append("Q3 宏F1 交叉核对不符")

    cfg = json.loads((BASE / "output/tables/q1_定稿配置.json").read_text(encoding="utf-8"))
    ts = cfg["theta_star"]
    ts_vals = list(ts.values()) if isinstance(ts, dict) else list(ts)
    lp = cfg["labels_props"]
    ok_q1 = (round(float(ts_vals[0])) == 5 and abs(float(ts_vals[1]) - 0.86) < 1e-9
             and abs(float(ts_vals[2]) - 0.98) < 1e-9
             and round(lp["合理诉求"], 4) == 0.8599 and round(lp["诉求偏高"], 4) == 0.1200
             and round(lp["严重超额"], 4) == 0.0201)
    emit("  Q1 交叉核对：theta_star=%s；占比 %.4f/%.4f/%.4f（期望 0.8599/0.1200/0.0201）-> %s"
         % (ts_vals, lp["合理诉求"], lp["诉求偏高"], lp["严重超额"], "PASS" if ok_q1 else "FAIL"))
    if not ok_q1:
        struct_fail.append("Q1 θ*/占比 交叉核对不符")

    # --- a6_01 日志重建确认 ---
    a601 = (BASE / "output/logs/a6_01_q1规则标注.log").stat().st_size
    emit("  a6_01 日志重建确认：size=%d 字节（>0 即已由主会话 17:59 重跑重建，WARN③ 处置落位）-> %s"
         % (a601, "PASS" if a601 > 0 else "FAIL"))
    if a601 == 0:
        struct_fail.append("a6_01 日志仍为 0 字节")

    # --- 渲染件与图表计数 ---
    n_pages = len(list((BASE / "output/render_pages").glob("page_*.png")))
    n_figs = len(list((BASE / "output/figures").glob("fig*.png")))
    emit("  渲染页数=%d（期望 25）、图 PNG 数=%d（期望 8）-> %s"
         % (n_pages, n_figs, "PASS" if (n_pages == 25 and n_figs == 8) else "FAIL"))
    if n_pages != 25 or n_figs != 8:
        struct_fail.append("渲染页/图数不符")
except Exception as e:
    struct_fail.append("结构复验异常: %r" % e)
    emit("  [异常] 结构复验失败：%r" % e)
emit("")

# ---------------------------------------------------------------- 目录余量扫描（防遗漏）
emit("## 目录余量扫描（期望清单之外的实际文件）")
EXPECT = set()
for _, rels in GROUPS:
    EXPECT.update(rels)
EXPECT.add("output/logs/" + SELF_LOG_NAME)
extras = []
for sub in ["code", "output/logs", "output/tables", "output/figures", "output/render_pages",
            "00_admin", "paper", "data"]:
    for p in (BASE / sub).rglob("*"):
        if p.is_file():
            rel = p.relative_to(BASE).as_posix()
            if "/__pycache__/" in "/" + rel or rel.startswith("code/__pycache__/"):
                continue
            if rel not in EXPECT:
                extras.append((rel, p.stat().st_size))
if extras:
    for rel, sz in sorted(extras):
        emit("  [EXTRA  ] %-70s size=%d" % (rel, sz))
else:
    emit("  无目录余量：工作区交付相关目录内文件均已列入期望清单。")
emit("")

# ---------------------------------------------------------------- 汇总判定
emit("## 汇总判定")
emit("  逐文件核对：期望 %d 项，OK %d，MISSING %d，EMPTY(0字节) %d"
     % (grand_total, grand_ok, len(missing), len(empty)))
for rel in missing:
    emit("    MISSING -> %s" % rel)
for rel in empty:
    emit("    EMPTY   -> %s" % rel)
emit("  结构复验失败项：%d %s" % (len(struct_fail), struct_fail if struct_fail else ""))
verdict = (not missing) and (not empty) and (not struct_fail)
emit("  总判定：%s" % ("PASS —— 交付包逐文件齐备，结构复验全过" if verdict else "FAIL —— 存在缺口，禁止打包放行"))
emit("  （G9 交付清单以本日志为准编制：output/交付清单.md）")

LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8")
print("\n[log written] %s (%d bytes)" % (LOG, LOG.stat().st_size))
sys.exit(0 if verdict else 1)
