# -*- coding: utf-8 -*-
"""G0 环境自检脚本 —— 盲测 Run-2（A0 阶段产物）

用途:
    核验 Python 环境与 data/ 四个输入文件的开放性。
    只记录"有什么可用"与文件结构信息（sheet 名、行列数），
    不做任何模型选型推荐，不深入分析数据内容（深入勘察是 A3 的职责）。

运行:
    python code/g0_env_check.py

输出:
    stdout + output/logs/g0_env_check.log（utf-8，供全阶段溯源引用）
"""
import sys
import os
import shutil
import importlib
import datetime

WS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOG = os.path.join(WS, "output", "logs", "g0_env_check.log")
os.makedirs(os.path.dirname(LOG), exist_ok=True)

lines = []


def out(s=""):
    lines.append(str(s))


out(f"# G0 环境自检  运行时间: {datetime.datetime.now().isoformat()}")
out(f"# 工作区: {WS}")
out(f"# Python: {sys.version.replace(chr(10), ' ')}")
out("")

# ---------- 1. 库可用性 ----------
out("## 1. 库可用性（逐个 import 测试）")
libs = ["numpy", "pandas", "sklearn", "scipy", "matplotlib",
        "statsmodels", "openpyxl", "docx", "lightgbm", "xgboost", "seaborn"]
lib_status = {}
for name in libs:
    try:
        m = importlib.import_module(name)
        ver = getattr(m, "__version__", "unknown")
        lib_status[name] = "OK"
        out(f"  [OK]   {name:<12s} version={ver}")
    except Exception as e:
        lib_status[name] = f"FAIL:{type(e).__name__}:{e}"
        out(f"  [FAIL] {name:<12s} {type(e).__name__}: {e}")
out("")

# ---------- 2. pandas + openpyxl 引擎往返 ----------
out("## 2. pandas 读取/写入引擎（openpyxl）往返测试")
try:
    import pandas as pd
    df = pd.DataFrame({"a": [1, 2]})
    tmp = os.path.join(os.path.dirname(LOG), "_engine_test.xlsx")
    with pd.ExcelWriter(tmp, engine="openpyxl") as w:
        df.to_excel(w, index=False)
    back = pd.read_excel(tmp, engine="openpyxl")
    ok = (list(back.columns) == ["a"]) and (len(back) == 2)
    os.remove(tmp)
    out(f"  pandas {pd.__version__} + openpyxl 引擎 写/读往返: {'OK' if ok else 'MISMATCH'}")
except Exception as e:
    out(f"  FAIL {type(e).__name__}: {e}")
out("")

# ---------- 3. data/ 文件开放性（仅结构） ----------
out("## 3. data/ 四文件存在性与 xlsx 开放性（仅报 sheet 名与行列数，不读内容）")
data_dir = os.path.join(WS, "data")
expected = ["赛题原文.md", "附件1.xlsx", "附件2.xlsx", "Result.xlsx"]
for f in expected:
    p = os.path.join(data_dir, f)
    if not os.path.exists(p):
        out(f"  [缺失] {f}")
        continue
    out(f"  [存在] {f}  size={os.path.getsize(p)} bytes")
    if f.endswith(".md"):
        try:
            with open(p, "r", encoding="utf-8") as fh:
                n_lines = sum(1 for _ in fh)
            out(f"      文本可读(utf-8), 行数={n_lines}")
        except Exception as e:
            out(f"      [读取失败] {type(e).__name__}: {e}")
    elif f.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True, data_only=True)
            out(f"      sheets={wb.sheetnames}")
            for ws in wb.worksheets:
                mr = ws.max_row if ws.max_row is not None else "未知"
                mc = ws.max_column if ws.max_column is not None else "未知"
                out(f"      sheet='{ws.title}'  max_row={mr}  max_col={mc}")
            wb.close()
        except Exception as e:
            out(f"      [openpyxl 打开失败] {type(e).__name__}: {e}")
out("")

# ---------- 4. matplotlib 中文字体 ----------
out("## 4. matplotlib 中文字体可用性")
try:
    import matplotlib
    from matplotlib import font_manager as fm
    names = sorted({f.name for f in fm.fontManager.ttflist})
    candidates = ["SimHei", "Microsoft YaHei", "SimSun", "KaiTi", "FangSong",
                  "DengXian", "Noto Sans CJK SC", "Source Han Sans SC",
                  "SourceHanSansSC", "STSong", "STKaiti", "STXihei", "LiSu",
                  "YouYuan", "WenQuanYi Zen Hei", "AR PL UMing CN"]
    found = [c for c in candidates if c in names]
    extra = [n for n in names
             if any(k in n for k in ("CJK", "Han", "Hei", "Song", "Kai",
                                     "Ming", "YaHei", "WenQuanYi"))
             and n not in found]
    out(f"  matplotlib version={matplotlib.__version__}, 已注册字体族数={len(names)}")
    out(f"  常用中文字体命中: {found if found else '无'}")
    out(f"  其他疑似中文字体（前10）: {extra[:10] if extra else '无'}")
except Exception as e:
    out(f"  FAIL {type(e).__name__}: {e}")
out("")

# ---------- 5. pandoc ----------
out("## 5. pandoc")
pandoc = shutil.which("pandoc")
out(f"  pandoc: {pandoc if pandoc else '未找到（PATH 中不存在）'}")
out("")

# ---------- 6. 工作区目录结构 ----------
out("## 6. 工作区目录结构（初始化后的约定子目录）")
for sub in ["00_admin", "code", "data", "output", "output/logs",
            "output/tables", "output/figures", "paper"]:
    d = os.path.join(WS, *sub.split("/"))
    if not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)
        out(f"  [新建] {sub}/")
    else:
        out(f"  [已有] {sub}/")

report = "\n".join(lines) + "\n"
with open(LOG, "w", encoding="utf-8") as fh:
    fh.write(report)
sys.stdout.reconfigure(encoding="utf-8")
print(report)
print(f"[日志已写入] {LOG}")
